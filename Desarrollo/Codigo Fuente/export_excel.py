"""Exportación del reporte Excel: todas las pestañas y subpestañas con la configuración base (sin filtros),
con las mismas tablas de la interfaz y gráficos nativos de Excel vinculados a las celdas (editables).

Cada hoja indica en su encabezado el nombre del escenario («Caso …») y, debajo, la fecha de la corrida.
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import fmt
import modelo
import reporte as rp

FUENTE = "Segoe UI"
FORMATOS = {"entero": fmt.XL_ENTERO, "dec1": fmt.XL_UN_DECIMAL, "pct": fmt.XL_PORCENTAJE,
            "dec2": "#,##0.00", "mt": "#,##0.00", "mt6": "#,##0.000000"}
_BORDE = Border(bottom=Side(style="thin", color="C9CFD6"))
SERIES = ["0969DA", "1A7F37", "BF8700", "8250DF", "BC4C00", "1B7C83", "CF222E", "6E7781", "4D2D00", "0550AE"]
COLOR_TIPO = {"Expit": rp.COLOR_ESTADO["TP"], "Rehandle": rp.COLOR_ESTADO["DPP"], "N/A": rp.COLOR_ESTADO["SB"],
              "Chancador": rp.COLOR_ESTADO["DENP"], "Stockpile": rp.COLOR_ESTADO["DEP"],
              "Botadero": rp.COLOR_ESTADO["DPNP"], "Total": "57606A"}
# Un color tenue por grupo (coherente con la interfaz) y un tono más marcado para subtotales
GRUPOS_XL = ["EEF4FC", "EDF7F0", "FCF7EA", "F3F0FC", "ECF6F7", "FCF1F0"]
GRUPOS_FUERTE_XL = ["DCE8F8", "DCEFE2", "F7ECCD", "E6E0F8", "D9EDEF", "F7E0DD"]


# ----------------------------------------------------------------------------------------
# Utilidades de formato
# ----------------------------------------------------------------------------------------
_ESTILOS = {}


def _cache(clave, fabrica):
    """Objetos de estilo compartidos (las tablas diarias tienen decenas de miles de celdas)."""
    v = _ESTILOS.get(clave)
    if v is None:
        v = _ESTILOS[clave] = fabrica()
    return v


def _relleno(hex_):
    return _cache(("fill", hex_), lambda: PatternFill("solid", start_color=hex_, end_color=hex_))


def _subtitulo(ws, texto, fila, col=1):
    c = ws.cell(row=fila, column=col, value=texto)
    c.font = Font(name=FUENTE, size=12, bold=True, color=rp.COLOR_TEXTO)


def _texto(ws, fila, col, texto, suave=False, negrita=False):
    c = ws.cell(row=fila, column=col, value=texto)
    c.font = Font(name=FUENTE, size=10, bold=negrita, color=rp.COLOR_SUAVE if suave else rp.COLOR_TEXTO)
    return c


def fecha_corrida(esc):
    try:
        return datetime.fromisoformat((esc.resultados or {}).get("fecha")).strftime("%H:%M %d/%m/%Y")
    except (TypeError, ValueError):
        return fmt.fecha_hora().replace(" ", "  ")


def _portada(ws, esc, titulo, detalle=""):
    """Encabezado de cada hoja: título, «Caso <escenario>» y, debajo, la fecha de la corrida."""
    ws.sheet_view.showGridLines = False
    c = ws.cell(row=1, column=1, value=titulo)
    c.font = Font(name=FUENTE, size=16, bold=True, color=rp.COLOR_ACENTO)
    c = ws.cell(row=2, column=1, value=f"Caso {esc.nombre_visible}")
    c.font = Font(name=FUENTE, size=11, bold=True, color="0000FF")
    _texto(ws, 3, 1, f"Fecha de Corrida: {fecha_corrida(esc)}" + (f"   ·   {detalle}" if detalle else ""), suave=True)
    return 5


def _tabla(ws, fila0, columnas, filas, totales=(), col0=1, grupos=(), subtotales=(), niveles=None):
    """Tabla con encabezado. tipo: 'texto', una clave de FORMATOS o un formato de Excel. 'subtotales': filas
    resaltadas; 'niveles': nivel de esquema por fila (para contraer en Excel). Devuelve la siguiente fila libre."""
    for j, (_k, titulo, tipo) in enumerate(columnas):
        c = ws.cell(row=fila0, column=col0 + j, value=titulo)
        c.font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
        c.fill = _relleno(rp.COLOR_CAB)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    lineas = max((str(t).count("\n") + 1 for _k, t, _ty in columnas), default=1)
    ws.row_dimensions[fila0].height = max(30, 15 * lineas + 6)
    gcol = _columna_grupo(columnas, filas, totales)
    grupo, previo = -1, object()
    for i, datos in enumerate(filas):
        fila = fila0 + 1 + i
        es_total, es_grupo, es_sub = i in totales, i in grupos, i in subtotales
        if gcol is not None and not es_total:
            clave = str(datos[gcol]).strip()
            if clave != previo:
                grupo, previo = grupo + 1, clave
            fondo = (GRUPOS_FUERTE_XL if (es_sub or es_grupo) else GRUPOS_XL)[grupo % len(GRUPOS_XL)]
        else:
            fondo = rp.COLOR_TOTAL if es_total else (GRUPOS_FUERTE_XL[0] if es_sub else (rp.COLOR_ALT if i % 2 == 1 else None))
        negrita = es_total or es_grupo or es_sub
        fuente = _cache(("font", negrita), lambda: Font(name=FUENTE, size=10, bold=negrita, color=rp.COLOR_TEXTO))
        izq = _cache(("al", "left"), lambda: Alignment(horizontal="left", indent=1))
        der = _cache(("al", "right"), lambda: Alignment(horizontal="right", indent=1))
        for j, ((_k, _t, tipo), v) in enumerate(zip(columnas, datos)):
            c = ws.cell(row=fila, column=col0 + j, value=v if v is not None else ("" if tipo == "texto" else "–"))
            c.font = fuente
            c.border = _BORDE
            if tipo == "texto":
                c.alignment = izq
            else:
                c.alignment = der
                if isinstance(v, (int, float)):
                    c.number_format = FORMATOS.get(tipo, tipo)
            if es_total:
                c.fill = _relleno(rp.COLOR_TOTAL)
            elif fondo:
                c.fill = _relleno(fondo)
        if niveles is not None and i < len(niveles) and niveles[i]:
            ws.row_dimensions[fila].outlineLevel = niveles[i]
    return fila0 + len(filas) + 2


def _columna_grupo(columnas, filas, totales):
    cuerpo = [f for i, f in enumerate(filas) if i not in totales]
    if len(cuerpo) < 3:
        return None
    for j, (_k, _t, tipo) in enumerate(columnas[:3]):
        if tipo != "texto":
            break
        vals = [str(f[j]).strip() for f in cuerpo if j < len(f)]
        if vals and all(vals) and len(set(vals)) < len(vals):
            return j
    return None


def _anchos(ws, minimo=11, maximo=34, primera=None):
    anchos = {}
    for fila in ws.iter_rows(min_row=5):
        for c in fila:
            if c.value is None:
                continue
            largo = len(str(c.value)) if not isinstance(c.value, float) else 11
            anchos[c.column] = max(anchos.get(c.column, 0), largo)
    for col, a in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = max(minimo, min(maximo, a * 1.12 + 3))
    if primera:
        ws.column_dimensions["A"].width = max(primera, ws.column_dimensions["A"].width or 0)


def _hoja(wb, nombre):
    return wb.create_sheet(nombre[:31])


# ----------------------------------------------------------------------------------------
# Gráficos nativos (vinculados a celdas)
# ----------------------------------------------------------------------------------------
def _grafico(hoja, ancla, titulo, cats, series, horizontal=True, agrupacion="clustered", alto=7.5, ancho=16.0,
             fmt_eje=None, etiquetas=True, leyenda=True, colores_puntos=None, fmt_etiquetas=None, umbral=None):
    """series: [(Reference datos, titulo, color|None)]; cats: Reference de categorías. 'umbral': las etiquetas
    de valores menores no se muestran (segmentos pequeños de barras apiladas)."""
    if umbral and fmt_etiquetas:
        fmt_etiquetas = f'[<{umbral:.6g}]"";{fmt_etiquetas}'
    g = BarChart()
    g.type = "bar" if horizontal else "col"
    g.grouping = agrupacion
    if agrupacion != "clustered":
        g.overlap = 100
    g.title = titulo
    g.title.overlay = False                     # el título no se dibuja sobre las barras
    g.height, g.width = alto, ancho
    g.x_axis.delete = False
    g.y_axis.delete = False
    g.y_axis.majorGridlines = None if horizontal else g.y_axis.majorGridlines
    if horizontal:
        g.x_axis.scaling.orientation = "maxMin"     # mismo orden que la tabla (eje de valores arriba)
    if fmt_eje:
        g.y_axis.numFmt = fmt_eje
    if agrupacion == "percentStacked":
        g.y_axis.scaling.max = 1
    for i, (ref, nombre, color) in enumerate(series):
        s = Series(ref, title=nombre)
        col = color or SERIES[i % len(SERIES)]
        if col == "ninguno":
            s.graphicalProperties.noFill = True
            s.graphicalProperties.line.noFill = True
        else:
            s.graphicalProperties.solidFill = col
            s.graphicalProperties.line.solidFill = col
        if etiquetas and col != "ninguno":
            s.dLbls = DataLabelList()
            s.dLbls.showVal = True
            for k in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
                setattr(s.dLbls, k, False)
            if fmt_etiquetas:
                s.dLbls.numFmt = fmt_etiquetas
        g.series.append(s)
    if colores_puntos:
        s = g.series[-1]
        for i, c in enumerate(colores_puntos):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = c
            pt.graphicalProperties.line.solidFill = c
            s.dPt.append(pt)
    g.set_categories(cats)
    if not leyenda:
        g.legend = None
    else:
        g.legend.position = "b"
        g.legend.overlay = False
    hoja.add_chart(g, ancla)
    return g


def _filas_grafico(alto_cm):
    return int(alto_cm / 0.5) + 2


# ----------------------------------------------------------------------------------------
# Reporte
# ----------------------------------------------------------------------------------------
def exportar_reporte(ruta, esc, progreso=None):
    """Libro con todas las pestañas y subpestañas (configuración base, sin filtros)."""
    if not esc.resultados or not esc.resultados.get("clases"):
        raise ValueError("No hay resultados para exportar.")
    wb = Workbook()
    pasos = [("Resumen", lambda: _hoja_resumen(wb.active, esc)),
             ("Análisis de Estados", lambda: _hojas_estados(wb, esc)),
             ("Análisis de Tiempos", lambda: [_hoja_tiempos(wb, esc, u) for u in ("dia", "anio")]),
             ("Orígenes y Destinos", lambda: _hoja_cargas(wb, esc)),
             ("Tiempos Ciclo", lambda: _hoja_tiempos_ciclo(wb, esc)),
             ("Métricas", lambda: _hoja_metricas(wb, esc)),
             ("Plan de Mina", lambda: _hojas_plan_mina(wb, esc)),
             ("Plan vs Simulación", lambda: _hojas_plan(wb, esc)),
             ("Tablas Dinámicas", lambda: (_hojas_pivot_estados(wb, esc), _hoja_pivot_registro(wb, esc))),
             ("Parámetros", lambda: _hoja_parametros(wb, esc))]
    for i, (nombre, fn) in enumerate(pasos):
        if progreso:
            progreso(i / (len(pasos) + 1), f"Hoja «{nombre}»…")
        fn()
    if progreso:
        progreso(0.97, "Guardando el libro…")
    guardar(wb, ruta)
    return ruta


def guardar(wb, ruta):
    """Guarda el libro y marca el formato de las etiquetas de datos como propio (sourceLinked="0"); sin esa
    marca Excel usa el formato de las celdas e ignora, por ejemplo, el umbral de las etiquetas pequeñas."""
    import os
    import re
    import shutil
    import tempfile
    import zipfile
    wb.save(ruta)
    patron = re.compile(r'<numFmt formatCode="([^"]*)"\s*/>')
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(os.path.abspath(ruta)))
    os.close(fd)
    try:
        with zipfile.ZipFile(ruta) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                datos = zin.read(item.filename)
                if item.filename.startswith("xl/charts/chart") and item.filename.endswith(".xml"):
                    datos = patron.sub(lambda m: f'<numFmt formatCode="{m.group(1)}" sourceLinked="0"/>',
                                       datos.decode("utf-8")).encode("utf-8")
                zout.writestr(item, datos)
        shutil.move(tmp, ruta)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def _hoja_resumen(ws, esc):
    res = esc.resultados
    ws.title = "Resumen"
    import config
    fila = _portada(ws, esc, f"{config.NOMBRE_PROGRAMA} – Resumen", f"Réplica: {res.get('replica', '')}")
    if esc.descripcion.strip():
        _texto(ws, fila, 1, "Descripción:", suave=True)
        c = _texto(ws, fila, 2, esc.descripcion.strip())
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=8)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[fila].height = max(30, 15 * (len(esc.descripcion) // 90 + 1))
        fila += 2
    _subtitulo(ws, "Indicadores por Flota", fila)
    fila += 1
    cols = [("clase", "Clase de Equipo", "texto"), ("tipo", "Tipo", "texto"), ("flota", "Flota", "texto"),
            ("n", "N° de Equipos", "entero"), ("hp", "Horas Productivas (h)", "entero"),
            ("disp", "Disponibilidad (%)", "pct"), ("util", "Utilización (%)", "pct"),
            ("mn", "Material / Taladros", "texto"), ("m", "Total", "dec1")]
    resumen = rp.resumen_flotas(res)
    filas = [[r["clase"], r["tipo"], r["flota"], r["n"], r["hp"], r["disp"], r["util"],
              "Material Total (Mt)" if r["material"] else "N° Taladros Total", r["metrica"]] for r in resumen]
    fila0 = fila
    fila = _tabla(ws, fila0, cols, filas)
    for i, r in enumerate(resumen):
        ws.cell(row=fila0 + 1 + i, column=9).number_format = fmt.XL_UN_DECIMAL if r["material"] else fmt.XL_ENTERO
    n = len(filas)
    if n:
        _grafico(ws, f"K{fila0}", "Disponibilidad y Utilización por Flota",
                 Reference(ws, min_col=3, min_row=fila0 + 1, max_row=fila0 + n),
                 [(Reference(ws, min_col=6, min_row=fila0 + 1, max_row=fila0 + n), "Disponibilidad (%)", SERIES[0]),
                  (Reference(ws, min_col=7, min_row=fila0 + 1, max_row=fila0 + n), "Utilización (%)", SERIES[1])],
                 alto=4 + 0.6 * n, ancho=17, fmt_eje="0%", fmt_etiquetas="0.0%")
    _anchos(ws, primera=18)


def _distribucion(hoja, fila, filas, titulo, col_nombre="Flota"):
    """Tabla de % del tiempo calendario por estado + gráfico 100 % apilado."""
    cols = [("f", col_nombre, "texto")] + [(k, modelo.NOMBRE_ESTADO[k], "pct") for k in modelo.CLAVES]
    datos = [[str(f["nombre"])] + [(f.get(k.lower(), 0) / f["tc"]) if f.get("tc") else 0 for k in modelo.CLAVES]
             for f in filas]
    ini = fila + 1
    _subtitulo(hoja, titulo, fila)
    fila = _tabla(hoja, ini, cols, datos)
    n = len(datos)
    if 0 < n <= 80:
        alto = 4.5 + 0.55 * n
        _grafico(hoja, f"J{ini}", titulo, Reference(hoja, min_col=1, min_row=ini + 1, max_row=ini + n),
                 [(Reference(hoja, min_col=2 + j, min_row=ini + 1, max_row=ini + n), modelo.NOMBRE_ESTADO[k],
                   rp.COLOR_ESTADO[k]) for j, k in enumerate(modelo.CLAVES)],
                 agrupacion="percentStacked", alto=alto, ancho=20, fmt_eje="0%", fmt_etiquetas="0.0%",
                 etiquetas=n <= 25, umbral=0.035)
        fila = max(fila, ini + _filas_grafico(alto))
    return fila


def _hojas_estados(wb, esc):
    """Análisis de Estados: por clase, vista Resumen (distribución, por flota/tipo/ID y distribución por ID)
    y Detallado (horas por estado) en la misma hoja."""
    res = esc.resultados
    for clase in modelo.ORDEN_CLASES:
        r = res["clases"].get(clase)
        if not r or not r["flota"]:
            continue
        info = modelo.CLASES[clase]
        hoja = _hoja(wb, f"Estados {info['plural']}")
        fila = _portada(hoja, esc, f"Análisis de Estados – {info['plural']}",
                        f"N° de {info['plural']}: {r['n_equipos']:,}   ·   Réplica: {r.get('replica', '')}")
        fila = _distribucion(hoja, fila, r["flota"], "Distribución del Tiempo Calendario (%)")
        for t in rp.tablas_clase(r, clase, detalle=True):
            _subtitulo(hoja, t["titulo"], fila)
            fila = _tabla(hoja, fila + 1, t["columnas"], t["filas"], t["totales"])
        fila = _distribucion(hoja, fila, r["id"], "Distribución por ID (%)", "ID")
        _anchos(hoja, primera=16)


def _hoja_tiempos(wb, esc, unidad):
    """Análisis de Tiempos (h/día-eq o h/periodo-eq): tablas por flota, resumen por estado y gráficos."""
    tiempos = (esc.resultados or {}).get("tiempos") or {}
    if not tiempos:
        return
    u = modelo.UNIDADES[unidad]
    formato = "#,##0.00" if unidad == "dia" else fmt.XL_UN_DECIMAL
    hoja = _hoja(wb, "Tiempos Diarios" if unidad == "dia" else "Tiempos Periodo")
    fila = _portada(hoja, esc, f"Análisis de Tiempos – {u['titulo']} ({u['unidad']})",
                    (f"Valor por equipo = Σ horas ÷ N° equipos ÷ {esc.dias:g} días" if unidad == "dia"
                     else "Valor por equipo = Σ horas ÷ N° equipos") + f"   ·   Réplica: {esc.resultados.get('replica', '')}")
    for clase in modelo.ORDEN_CLASES:
        t = tiempos.get(clase)
        if not t or not t.get("ids"):
            continue
        info = modelo.CLASES[clase]
        _subtitulo(hoja, info["plural"], fila)
        fila += 2
        resumen, actividades = [], []
        for flota in modelo.flotas_de(t):
            ids = [e for e in t["ids"] if e["flota"] == flota]
            b = modelo.bloque_tiempos(t, ids, unidad, esc.dias)
            if not b:
                continue
            resumen.append((flota, b))
            _texto(hoja, fila, 1, f"{flota}  ·  {b['n']} und", negrita=True)
            fila += 1
            cols = [("e", "Estado", "texto")] + [(f"a{j}", c["nombre"].replace("(hs)", "").strip(), formato)
                                                  for j, c in enumerate(t["columnas"])] + [("h", f"Horas ({u['unidad']})", formato)]
            filas = [[modelo.FILA_TIEMPOS[k]] + [v if c["estado"] == k else "" for v, c in zip(b["valores"], t["columnas"])]
                     + [b["por_estado"][k]] for k in modelo.ORDEN_TIEMPOS]
            filas.append(["Total"] + list(b["valores"]) + [b["total"]])
            ini = fila
            fin = _tabla(hoja, fila, cols, filas, totales=[len(filas) - 1])
            actividades.append((flota, b, ini, len(t["columnas"]), t["columnas"]))
            c0 = len(cols) + 2
            for i, (et, val, f_) in enumerate((("Horas Productivas", b["horas_productivas"], formato),
                                               ("Disponibilidad", b["disponibilidad"], fmt.XL_PORCENTAJE),
                                               ("Utilización", b["utilizacion"], fmt.XL_PORCENTAJE),
                                               (u["total"], b["total"], formato))):
                _texto(hoja, ini + 1 + i, c0, et, suave=True)
                celda = hoja.cell(row=ini + 1 + i, column=c0 + 1, value=val)
                celda.number_format = f_
                celda.font = Font(name=FUENTE, size=10, bold=True, color=rp.COLOR_TEXTO)
            fila = fin
        if not resumen:
            continue
        _texto(hoja, fila, 1, f"Distribución por Estado – {info['plural']} ({u['unidad']})", negrita=True)
        fila += 1
        cols_r = [("f", "Flota", "texto"), ("n", "Und", "entero")] + \
                 [(k, modelo.FILA_TIEMPOS[k], formato) for k in modelo.ORDEN_TIEMPOS] + [("t", "Total", formato)]
        datos_r = [[f, b["n"]] + [b["por_estado"][k] for k in modelo.ORDEN_TIEMPOS] + [b["total"]] for f, b in resumen]
        ini = fila
        fila = _tabla(hoja, fila, cols_r, datos_r)
        n = len(resumen)
        alto = 4.5 + 0.8 * n
        g = _grafico(hoja, f"A{fila}", f"Distribución por Estado – {info['plural']} ({u['unidad']})",
                     Reference(hoja, min_col=1, min_row=ini + 1, max_row=ini + n),
                     [(Reference(hoja, min_col=3 + j, min_row=ini + 1, max_row=ini + n), modelo.FILA_TIEMPOS[k],
                       rp.COLOR_ESTADO[k]) for j, k in enumerate(modelo.ORDEN_TIEMPOS)],
                     agrupacion="stacked", alto=alto, ancho=22, fmt_eje="#,##0", fmt_etiquetas="#,##0.0",
                     umbral=0.045 * max(b["total"] for _f, b in resumen))
        g.y_axis.majorUnit = modelo.paso_eje(unidad, esc.dias)
        g.y_axis.scaling.min = 0
        if max(b["total"] for _f, b in resumen) <= modelo.paso_eje(unidad, esc.dias) * 6 * 1.0005:
            g.y_axis.scaling.max = modelo.paso_eje(unidad, esc.dias) * 6
        fila += _filas_grafico(alto)
        for k, (flota, b, ini_b, n_act, cols_t) in enumerate(actividades):
            alto_a = 3.5 + 0.5 * n_act
            _grafico(hoja, f"{'A' if k % 2 == 0 else 'J'}{fila}", f"Horas por Actividad – {flota} ({u['unidad']})",
                     Reference(hoja, min_col=2, max_col=1 + n_act, min_row=ini_b),
                     [(Reference(hoja, min_col=2, max_col=1 + n_act, min_row=ini_b + 7), f"Horas ({u['unidad']})", None)],
                     alto=alto_a, ancho=15, leyenda=False, fmt_etiquetas="#,##0.00" if unidad == "dia" else "#,##0",
                     colores_puntos=[rp.COLOR_ESTADO[c["estado"]] for c in cols_t])
            if k % 2 == 1 or k == len(actividades) - 1:
                fila += _filas_grafico(alto_a)
        fila += 1
    _anchos(hoja, maximo=20, primera=24)


def _hoja_cargas(wb, esc):
    """Orígenes y Destinos: cascadas por tipo de origen y de destino, desglose por fase y polígono y flujo
    Tipo de Origen ▸ Condición ▸ Polígono → Tipo de Destino."""
    import registro
    cargas = (esc.resultados or {}).get("cargas")
    if not cargas:
        return
    hoja = _hoja(wb, "Orígenes y Destinos")
    fila = _portada(hoja, esc, "Orígenes y Destinos", f"Réplica: {esc.resultados.get('replica', '')}   ·   Material en Mt")
    fases = registro.fases_por_tipo(cargas, esc.origenes)
    for titulo, datos, tipos, col in (
            ("Reporte por Tipo de Origen", registro.por_tipo_origen(cargas, esc.origenes), registro.TIPOS_ORIGEN,
             "Tipo de Origen"),
            ("Reporte por Tipo de Destino", registro.por_tipo_destino(cargas, esc.destinos), registro.TIPOS_DESTINO,
             "Tipo de Destino")):
        _subtitulo(hoja, titulo, fila)
        fila += 1
        tipos = [t for t in tipos if datos["tipos"].get(t, 0) > 0 or t != "N/A"]
        total = datos["total"] or 1.0
        filas, acum = [], 0.0
        for t in tipos:
            v = datos["tipos"].get(t, 0.0) / 1e6
            filas.append([t, v, v / (total / 1e6), acum])
            acum += v
        filas.append(["Total", datos["total"] / 1e6, 1.0 if datos["total"] else None, 0.0])
        ini = fila
        fila = _tabla(hoja, fila, [("t", col, "texto"), ("m", "Material (Mt)", "mt"), ("p", "Participación (%)", "pct"),
                                   ("b", "Base Cascada (Mt)", "mt")], filas, totales=[len(filas) - 1])
        n = len(filas)
        _grafico(hoja, f"F{ini}", titulo, Reference(hoja, min_col=1, min_row=ini + 1, max_row=ini + n),
                 [(Reference(hoja, min_col=4, min_row=ini + 1, max_row=ini + n), "Base", "ninguno"),
                  (Reference(hoja, min_col=2, min_row=ini + 1, max_row=ini + n), "Material (Mt)", None)],
                 horizontal=False, agrupacion="stacked", alto=7, ancho=13, leyenda=False, fmt_eje="#,##0",
                 fmt_etiquetas="#,##0.0", colores_puntos=[COLOR_TIPO.get(t, SERIES[0]) for t in tipos + ["Total"]])
        fila = max(fila, ini + _filas_grafico(7))
        if col == "Tipo de Origen":
            _texto(hoja, fila, 1, "Detalle – Tipo de Origen ▸ Fase ▸ Polígono", negrita=True)
            fila += 1
            det, grupos, niveles = [], [], []
            for t in tipos:
                for fase, tf, polis in fases.get(t, []):
                    det.append([t, fase, "", tf / 1e6, tf / total])
                    grupos.append(len(det) - 1)
                    niveles.append(0)
                    for o, v in polis:
                        det.append([t, fase, o, v / 1e6, v / total])
                        niveles.append(1)
            fila = _tabla(hoja, fila, [("t", "Tipo de Origen", "texto"), ("f", "Fase", "texto"), ("o", "Polígono", "texto"),
                                       ("m", "Material (Mt)", "mt"), ("p", "Participación (%)", "pct")], det,
                          grupos=grupos, niveles=niveles)
        else:
            _texto(hoja, fila, 1, "Detalle – Tipo de Destino ▸ Destino", negrita=True)
            fila += 1
            det = [[t, nom, v / 1e6, v / total] for t in tipos
                   for nom, v in sorted(datos["detalle"].get(t, []), key=lambda x: x[0].lower())]
            fila = _tabla(hoja, fila, [("t", col, "texto"), ("n", "Destino", "texto"),
                                       ("m", "Material (Mt)", "mt"), ("p", "Participación (%)", "pct")], det)
    m = registro.flujo_poligonos(cargas, esc.origenes, esc.destinos)
    if m:
        _subtitulo(hoja, "Flujo de Material: Tipo de Origen ▸ Condición ▸ Polígono → Tipo de Destino (Mt)", fila)
        fila += 1
        filas, grupos, niveles = [], [], []
        orden = sorted(m.items(), key=lambda x: (registro.TIPOS_ORIGEN.index(x[0][0]) if x[0][0] in registro.TIPOS_ORIGEN
                                                 else 9, registro.CONDICIONES.index(x[0][1])
                                                 if x[0][1] in registro.CONDICIONES else 9))
        tot = [0.0] * (len(registro.TIPOS_DESTINO) + 1)

        def fila_v(cnt):
            s = sum(cnt.get(t, 0.0) for t in registro.TIPOS_DESTINO)
            return [cnt.get(t, 0.0) / 1e6 for t in registro.TIPOS_DESTINO] + [s / 1e6, cnt.get("Chancador", 0.0) / s if s else None]
        for (t_o, cond), polis in orden:
            suma = {t: sum(p.get(t, 0.0) for p in polis.values()) for t in registro.TIPOS_DESTINO}
            for i, t in enumerate(registro.TIPOS_DESTINO):
                tot[i] += suma[t]
            for o, cnt in sorted(polis.items(), key=lambda x: -sum(x[1].values())):
                filas.append([t_o, cond, o] + fila_v(cnt))
                niveles.append(1)
            filas.append([t_o, cond, "Subtotal"] + fila_v(suma))
            grupos.append(len(filas) - 1)
            niveles.append(0)
        tot[-1] = sum(tot[:-1])
        filas.append(["Total", "", ""] + [v / 1e6 for v in tot] + [tot[0] / tot[-1] if tot[-1] else None])
        niveles.append(0)
        fila = _tabla(hoja, fila, [("o", "Tipo de Origen", "texto"), ("c", "Condición", "texto"),
                                   ("p", "Polígono", "texto")] +
                      [(t, f"{t} (Mt)", "mt") for t in registro.TIPOS_DESTINO] +
                      [("t", "Total (Mt)", "mt"), ("x", "Mineral a Planta (%)", "pct")], filas,
                      totales=[len(filas) - 1], subtotales=grupos, niveles=niveles)
    hoja.sheet_properties.outlinePr.summaryBelow = True
    _anchos(hoja, maximo=30, primera=20)


def _bloque_estadistica(hoja, fila, filas, titulo, unidad, decimales=2):
    """Tabla Mínimo · Promedio · Máximo por Flota de Pala y Tipo de Camión + gráfico (tres series)."""
    if not filas:
        return fila
    formato = "#,##0." + "0" * decimales if decimales else "#,##0"
    _subtitulo(hoja, f"{titulo} ({unidad})", fila)
    ini = fila + 1
    fila = _tabla(hoja, ini, [("p", "Flota Equipo Pala", "texto"), ("c", "Tipo Camiones", "texto"),
                              ("mn", f"Mínimo ({unidad})", formato), ("pr", f"Promedio ({unidad})", formato),
                              ("mx", f"Máximo ({unidad})", formato)],
                  [[f["pala"], f["camion"], f["min"], f["prom"], f["max"]] for f in filas])
    n = len(filas)
    for i, f in enumerate(filas):
        hoja.cell(row=ini + 1 + i, column=8, value=f"{f['pala']} – {f['camion']}")
    alto = 3.5 + 0.9 * n
    _grafico(hoja, f"J{ini}", f"{titulo} ({unidad})", Reference(hoja, min_col=8, min_row=ini + 1, max_row=ini + n),
             [(Reference(hoja, min_col=3 + j, min_row=ini + 1, max_row=ini + n), nombre, color)
              for j, (nombre, color) in enumerate((("Mínimo", "8C959F"), ("Promedio", "12355B"), ("Máximo", "3B8ED8")))],
             alto=alto, ancho=15, fmt_etiquetas=formato)
    return max(fila, ini + _filas_grafico(alto))


def _hoja_tiempos_ciclo(wb, esc):
    """Tiempos y Métricas ▸ Tiempos Ciclo: estadística por flota de pala y tipo de camión + Descarga Plan."""
    import registro
    res = esc.resultados or {}
    cargas = res.get("cargas")
    tp = (esc.clases.get("palas") or {}).get("tipos") or {}
    tc = (esc.clases.get("camiones") or {}).get("tipos") or {}
    if not cargas or not tp or not tc:
        return
    hoja = _hoja(wb, "Tiempos Ciclo")
    fila = _portada(hoja, esc, "Tiempos y Métricas – Tiempos Ciclo (min)",
                    f"Réplica: {res.get('replica', '')}   ·   Registros con carga; promedio por registro")
    if registro.tiene_estadisticas(cargas):
        medidas = {m: t for m, t, _u in registro.MEDIDAS_TIEMPO}
        for med in ("cg", "cu", "hang", "cc", "cola", "desc"):
            filas = registro.tabla_estadistica(cargas, med, tp, tc, esc.clases.get("palas"), esc.clases.get("camiones"))
            fila = _bloque_estadistica(hoja, fila, filas, medidas[med], "min")
    filas = [[f["tipo_destino"], f["camion"], f["cola"], f["descarga"], f["total"]]
             for f in registro.tabla_descarga(cargas, esc.destinos, tc)]
    if filas:
        _subtitulo(hoja, "Tiempo de Descarga Plan (min)", fila)
        ini = fila + 1
        fila = _tabla(hoja, ini, [("d", "Tipo de Destino", "texto"), ("c", "Tipo de Camión", "texto"),
                                  ("q", "Promedio Cola en Descarga (min)", "dec2"), ("x", "Promedio Descarga (min)", "dec2"),
                                  ("t", "Tiempo de Descarga Total (min)", "dec2")], filas)
        n = len(filas)
        for i, f in enumerate(filas):
            hoja.cell(row=ini + 1 + i, column=8, value=f"{f[0]} – {f[1]}")
        alto = 3.5 + 0.6 * n
        _grafico(hoja, f"J{ini}", "Tiempo de Descarga Plan (min)", Reference(hoja, min_col=8, min_row=ini + 1, max_row=ini + n),
                 [(Reference(hoja, min_col=3, min_row=ini + 1, max_row=ini + n), "Promedio Cola en Descarga",
                   rp.COLOR_ESTADO["DPNP"]),
                  (Reference(hoja, min_col=4, min_row=ini + 1, max_row=ini + n), "Promedio Descarga", rp.COLOR_ESTADO["TP"])],
                 agrupacion="stacked", alto=alto, ancho=15, fmt_etiquetas="0.00",
                 umbral=0.06 * max(f[4] or 0 for f in filas))
        fila = max(fila, ini + _filas_grafico(alto))
    _anchos(hoja, maximo=30, primera=18)


def _hoja_metricas(wb, esc):
    """Tiempos y Métricas ▸ Métricas: Registro de Ciclos (N° Ciclos, Ciclos (%) y Tonelaje (Mt)), N° de Pases,
    Carga por Camión, Productividad de Palas y de Camiones (Por Flota y Por ID) y N° Taladros."""
    import registro
    res = esc.resultados or {}
    clases = res.get("clases") or {}
    cargas = res.get("cargas")
    tp = (esc.clases.get("palas") or {}).get("tipos") or {}
    tc = (esc.clases.get("camiones") or {}).get("tipos") or {}
    cfg_p, cfg_c = esc.clases.get("palas"), esc.clases.get("camiones")
    hoja = _hoja(wb, "Métricas")
    fila = _portada(hoja, esc, "Tiempos y Métricas – Métricas", f"Réplica: {res.get('replica', '')}")
    if cargas and tp and tc and registro.tiene_estadisticas(cargas):
        r = registro.tabla_registro(cargas, tp, tc, cfg_p, cfg_c)
        palas, cams = r["palas"], r["camiones"]
        for modo, etq, tipo, div in (("n", "N° Ciclos", "entero", None), ("pct", "Ciclos (%)", "pct", "total"),
                                     ("t", "Tonelaje (Mt)", "mt", 1e6)):
            base = r["t"] if modo == "t" else r["n"]
            total = sum(base.get((p, c), 0) for p in palas for c in cams) or 1.0
            d = total if div == "total" else (div or 1.0)
            _subtitulo(hoja, f"Registro de Ciclos – {etq}", fila)
            ini = fila + 1
            filas = [[p] + [base.get((p, c), 0) / d for c in cams] + [sum(base.get((p, c), 0) for c in cams) / d]
                     for p in palas]
            filas.append(["Total"] + [sum(base.get((p, c), 0) for p in palas) / d for c in cams] + [total / d])
            fila = _tabla(hoja, ini, [("p", "Flota Equipo Pala", "texto")] + [(c, c, tipo) for c in cams] +
                          [("t", "Total", tipo)], filas, totales=[len(filas) - 1])
            if modo == "n" and palas and cams:
                alto = 3.5 + 0.7 * len(palas)
                _grafico(hoja, f"{get_column_letter(len(cams) + 4)}{ini}", "Registro de Ciclos – N° Ciclos",
                         Reference(hoja, min_col=1, min_row=ini + 1, max_row=ini + len(palas)),
                         [(Reference(hoja, min_col=2 + j, min_row=ini + 1, max_row=ini + len(palas)), c,
                           SERIES[0] if c == "UltraClass" else SERIES[2]) for j, c in enumerate(cams)],
                         agrupacion="stacked", alto=alto, ancho=15, fmt_eje="#,##0", fmt_etiquetas="#,##0",
                         umbral=0.07 * max(f[-1] for f in filas[:-1]))
                fila = max(fila, ini + _filas_grafico(alto))
        fila = _bloque_estadistica(hoja, fila, registro.tabla_estadistica(cargas, "ps", tp, tc, cfg_p, cfg_c),
                                   "N° de Pases", "pases")
        fila = _bloque_estadistica(hoja, fila, registro.tabla_estadistica(cargas, "t", tp, tc, cfg_p, cfg_c),
                                   "Carga por Camión", "t", decimales=1)

    def productividad(r, titulo, clase):
        nonlocal fila

        def prod(f):
            return f["metrica_total"] * 1e6 / f["horas_productivas"] if f.get("horas_productivas") else None
        for nivel in ("flota", "id"):
            _subtitulo(hoja, f"{titulo} – {'Por Flota' if nivel == 'flota' else 'Por ID'}", fila)
            ini = fila + 1
            if nivel == "flota":
                filas = [[f.get("tipo", ""), f["nombre"], f["horas_productivas"], f["metrica_total"], prod(f)]
                         for f in r["flota"]]
                cols = [("t", "Tipo", "texto"), ("f", "Flota", "texto")]
            else:
                filas = [[e.get("tipo", ""), e["flota"], str(e["nombre"]), e["horas_productivas"], e["metrica_total"],
                          prod(e)] for e in r["id"]]
                cols = [("t", "Tipo", "texto"), ("f", "Flota", "texto"), ("i", "ID", "texto")]
            cols += [("h", "Horas Productivas (h)", "entero"), ("m", "Material Movido (Mt)", "mt"),
                     ("p", "Productividad (t/h)", "entero")]
            fila = _tabla(hoja, ini, cols, filas)
            n = len(filas)
            if 0 < n <= 60:
                c_cat = 2 if nivel == "flota" else 3
                alto = 3.5 + 0.5 * n
                _grafico(hoja, f"I{ini}", f"{titulo} – {'Por Flota' if nivel == 'flota' else 'Por ID'}",
                         Reference(hoja, min_col=c_cat, min_row=ini + 1, max_row=ini + n),
                         [(Reference(hoja, min_col=len(cols), min_row=ini + 1, max_row=ini + n), "Productividad (t/h)", None)],
                         alto=alto, ancho=15, leyenda=False, fmt_etiquetas="#,##0")
                fila = max(fila, ini + _filas_grafico(alto))
    if clases.get("palas"):
        productividad(clases["palas"], "Productividad de Palas (t/h)", "palas")
    if clases.get("camiones"):
        productividad(clases["camiones"], "Productividad de Camiones (t/h)", "camiones")
    perf = clases.get("perforadoras")
    if perf:
        _subtitulo(hoja, "N° Taladros", fila)
        cat = (esc.clases.get("perforadoras") or {}).get("energia") or {}
        filas = [[f.get("tipo") or modelo.SIN_ASIGNAR, f["nombre"], cat.get(f["nombre"], ""), f["n_equipos"],
                  f["metrica_total"]] for f in perf["flota"]]
        filas.append(["Total", "", "", perf["n_equipos"], sum(f["metrica_total"] for f in perf["flota"])])
        fila = _tabla(hoja, fila + 1, [("t", "Tipo de Perforadora", "texto"), ("f", "Flota", "texto"),
                                       ("c", modelo.CATEGORIA, "texto"), ("n", "N° de Perforadoras", "entero"),
                                       ("x", "N° Taladros", "entero")], filas, totales=[len(filas) - 1])
    _anchos(hoja, maximo=30, primera=22)


def _hojas_plan_mina(wb, esc):
    """Plan de Mina por día, semana y mes (sin filtros): indicadores, tonelaje por flota de pala con gráfico,
    tonelaje por flota/ID y la tabla Fase ▸ Flota Pala ▸ ID Pala con subtotales (contraíble en Excel)."""
    import registro
    from vista_perfil import ADJETIVO, orden_flotas
    cargas = (esc.resultados or {}).get("cargas")
    if not cargas or not cargas.get("perfil"):
        return
    anio = esc.anio or 2000
    for modo in ("mes", "semana", "dia"):
        tabla = registro.tabla_perfil(cargas, anio, modo, esc.origenes, esc.destinos, {}, orden_flotas(esc))
        if not tabla["filas"]:
            continue
        nom = registro.NOMBRE_PERIODO.get(modo, "Mes")
        adj = ADJETIVO[modo]
        hoja = _hoja(wb, f"Plan de Mina {adj}")
        fila = _portada(hoja, esc, f"Plan de Mina – {adj} (t)",
                        f"Año {anio}   ·   Réplica: {esc.resultados.get('replica', '')}")
        per, tot = tabla["periodos"], tabla["totales"]
        con = [i for i in range(len(tot)) if tot[i] > 0] or [0]
        prom = tabla["total"] / len(con)
        i_max, i_min = max(con, key=lambda i: tot[i]), min(con, key=lambda i: tot[i])
        for j, (et, val, det) in enumerate((("Tonelaje Movido Total (t)", tabla["total"], ""),
                                            (f"Tonelaje Promedio {adj} (t)", prom, ""),
                                            (f"Tonelaje Máximo {adj} (t)", tot[i_max], per[i_max]),
                                            (f"Tonelaje Mínimo {adj} (t)", tot[i_min], per[i_min]))):
            _texto(hoja, fila, 1 + j * 3, et, suave=True)
            c = hoja.cell(row=fila + 1, column=1 + j * 3, value=val)
            c.number_format = fmt.XL_ENTERO
            c.font = Font(name=FUENTE, size=13, bold=True, color=rp.COLOR_TEXTO)
            if det:
                _texto(hoja, fila + 1, 2 + j * 3, det, suave=True)
        fila += 3
        flotas = [f for f in orden_flotas(esc) if any(x["flota"] == f for x in tabla["filas"])]
        flotas += [f for f in dict.fromkeys(x["flota"] for x in tabla["filas"]) if f not in flotas]
        _subtitulo(hoja, f"Tonelaje por {nom} y Flota Pala (t)", fila)
        ini = fila + 1
        cols = [("f", "Flota Pala", "texto")] + [(f"p{i}", p, "entero") for i, p in enumerate(per)] + [("t", "Total", "entero")]
        datos = []
        for fl in flotas:
            vals = [sum(f["valores"][i] for f in tabla["filas"] if f["flota"] == fl) for i in range(len(per))]
            datos.append([fl] + vals + [sum(vals)])
        datos.append(["Total"] + tot + [tabla["total"]])
        fila = _tabla(hoja, ini, cols, datos, totales=[len(datos) - 1])
        alto = 9
        _grafico(hoja, f"A{fila}", f"Tonelaje por {nom} y Flota Pala (t)",
                 Reference(hoja, min_col=2, max_col=1 + len(per), min_row=ini),
                 [(Reference(hoja, min_col=2, max_col=1 + len(per), min_row=ini + 1 + i), fl, SERIES[i % len(SERIES)])
                  for i, fl in enumerate(flotas)],
                 horizontal=False, agrupacion="stacked", alto=alto,
                 ancho=30 if modo == "mes" else min(40, 14 + 0.12 * len(per) * 4),
                 fmt_eje="#,##0", etiquetas=modo == "mes", fmt_etiquetas="#,##0,\"k\"",
                 umbral=0.06 * max(tot + [0]))
        fila += _filas_grafico(alto)
        n_per = len(per) or 1
        _subtitulo(hoja, f"Tonelaje por Flota Pala e ID Pala (t) – Total y Promedio {adj}", fila)
        filas, grupos, niveles = [], [], []
        for fl in flotas:
            ids = {}
            for f in tabla["filas"]:
                if f["flota"] == fl:
                    ids[f["id"]] = ids.get(f["id"], 0.0) + f["total"]
            s = sum(ids.values())
            for i, v in sorted(ids.items(), key=lambda x: -x[1]):
                filas.append([fl, i, v, v / (tabla["total"] or 1.0), v / n_per])
                niveles.append(1)
            filas.append([fl, f"{fl} Total", s, s / (tabla["total"] or 1.0), s / n_per])
            grupos.append(len(filas) - 1)
            niveles.append(0)
        filas.append(["Total", "", tabla["total"], 1.0, tabla["total"] / n_per])
        niveles.append(0)
        fila = _tabla(hoja, fila + 1, [("f", "Flota Pala", "texto"), ("i", "ID Pala", "texto"), ("t", "Tonelaje (t)", "entero"),
                                       ("p", "Participación (%)", "pct"), ("m", f"Promedio {adj} (t)", "entero")],
                      filas, totales=[len(filas) - 1], subtotales=grupos, niveles=niveles)
        # Fase ▸ Flota Pala ▸ ID Pala con subtotales (niveles de esquema de Excel para contraer)
        _subtitulo(hoja, "Tonelaje por Fase, Flota e ID de Pala (t)", fila)
        cols = [("f", "Fase", "texto"), ("l", "Flota Pala", "texto"), ("i", "ID Pala", "texto")] + \
               [(f"p{i}", p, "entero") for i, p in enumerate(per)] + [("t", "Total", "entero")]
        filas, subt, niveles = [], [], []
        fases = list(dict.fromkeys(f["fase"] for f in tabla["filas"]))
        for fase in fases:
            filas_f = [f for f in tabla["filas"] if f["fase"] == fase]
            for fl in list(dict.fromkeys(f["flota"] for f in filas_f)):
                filas_l = [f for f in filas_f if f["flota"] == fl]
                for f in filas_l:
                    filas.append([fase, fl, f["id"]] + [v or None for v in f["valores"]] + [f["total"]])
                    niveles.append(2)
                vals = [sum(f["valores"][i] for f in filas_l) for i in range(len(per))]
                filas.append([fase, f"{fl} Total", ""] + vals + [sum(vals)])
                subt.append(len(filas) - 1)
                niveles.append(1)
            vals = [sum(f["valores"][i] for f in filas_f) for i in range(len(per))]
            filas.append([f"{fase} Total", "", ""] + vals + [sum(vals)])
            subt.append(len(filas) - 1)
            niveles.append(0)
        filas.append(["Total", "", ""] + tot + [tabla["total"]])
        niveles.append(0)
        filas = [[x if x is not None else "" for x in r] for r in filas]
        fila = _tabla(hoja, fila + 1, cols, filas, totales=[len(filas) - 1], subtotales=subt, niveles=niveles)
        hoja.sheet_properties.outlinePr.summaryBelow = True
        hoja.freeze_panes = hoja.cell(row=5, column=2)
        _anchos(hoja, minimo=9, maximo=18, primera=14)


def _hojas_plan(wb, esc):
    import export_plan
    import plan
    if not esc.plan or not plan.replicas_resultados(esc):
        return
    ws = _hoja(wb, "Plan vs Simulación")
    export_plan.escribir(wb, ws, esc, "Gráficos vs Plan")


def _tabla_pivot(hoja, fila, tabla, cfg, titulo):
    """Escribe el resultado de un pivote y su gráfico (series y tipo como en la interfaz)."""
    import pivot
    n_dim = tabla["n_dim"]
    formatos = tabla.get("formatos") or [None] * len(tabla["columnas"])
    mapa = {"pct": "pct", "int": "entero", "0.00%": "pct", "0.0%": "0.0%", "#,##0": "entero", "#,##0.0": "dec1",
            "#,##0.00": "dec2", "#,##0.000": "#,##0.000"}
    cols = []
    for j, t in enumerate(tabla["columnas"]):
        if j < n_dim:
            cols.append((f"c{j}", t, "texto"))
            continue
        tipo = mapa.get(formatos[j] if j < len(formatos) else None, "dec1")
        vals = [f[j] for f in tabla["filas"] if isinstance(f[j], float)]
        if tipo == "dec1" and ("(h)" in t or (vals and all(abs(v - round(v)) < 1e-9 for v in vals))):
            tipo = "entero"
        cols.append((f"c{j}", t, tipo))
    ini = fila
    fila = _tabla(hoja, ini, cols, [[v for v in f] for f in tabla["filas"]], tabla["totales"])
    cuerpo = [i for i in range(len(tabla["filas"])) if i not in tabla["totales"]]
    if cfg.get("grafico", True) and cuerpo:
        series_cfg = cfg.get("grafico_series")
        idx = [j for j, c in enumerate(tabla["columnas"]) if j >= n_dim and pivot.TOTAL not in c
               and (not series_cfg or c in series_cfg)]
        if idx:
            base = "pct" if (formatos[idx[0]] if idx[0] < len(formatos) else None) in ("pct", "0.00%", "0.0%") else "num"
            idx = [j for j in idx if ("pct" if formatos[j] in ("pct", "0.00%", "0.0%") else "num") == base][:12]
            primero, ultimo = ini + 1 + cuerpo[0], ini + 1 + cuerpo[-1]
            n = ultimo - primero + 1
            horizontal = bool(cfg.get("grafico_horizontal"))
            alto = 7.5 if not horizontal else 4 + 0.6 * n
            _grafico(hoja, f"{get_column_letter(len(cols) + 2)}{ini}", titulo,
                     Reference(hoja, min_col=1, max_col=max(1, n_dim), min_row=primero, max_row=ultimo),
                     [(Reference(hoja, min_col=j + 1, min_row=primero, max_row=ultimo),
                       str(tabla["columnas"][j]).replace("\n", " · "), None) for j in idx],
                     horizontal=horizontal, agrupacion="stacked" if cfg.get("grafico_apilado") else "clustered",
                     alto=alto, ancho=18, fmt_eje="0%" if base == "pct" else "#,##0",
                     fmt_etiquetas="0.0%" if base == "pct" else "#,##0", etiquetas=n <= 25)
            fila = max(fila, ini + _filas_grafico(alto))
    return fila


def _hojas_pivot_estados(wb, esc):
    import copy
    import pivot
    from vista_pivot_estados import config_inicial
    res = esc.resultados or {}
    reps = [str(r) for r in res.get("replicas", []) if r != "Todas"]
    for clase in modelo.ORDEN_CLASES:
        if clase not in res.get("clases", {}):
            continue
        df = modelo.dataset_estados(res, clase, esc.clases.get(clase))
        if df.empty:
            continue
        guardado = (esc.pivots.get(clase) or {}).get("config")
        cfg = dict(config_inicial(clase, reps))
        if guardado:
            cfg.update(guardado)
        try:
            tabla = pivot.calcular(df, copy.deepcopy(cfg), modelo.MEDIDAS_ESTADOS)
        except Exception:
            continue
        info = modelo.CLASES[clase]
        hoja = _hoja(wb, f"Pivot Estados {info['plural']}")
        fila = _portada(hoja, esc, f"Análisis de Estados (Pivot) – {info['plural']}", _resumen_cfg(cfg))
        _tabla_pivot(hoja, fila, tabla, cfg, f"Análisis de Estados – {info['plural']}")
        _anchos(hoja, maximo=26)


def _hoja_pivot_registro(wb, esc):
    guardado = esc.pivot if isinstance(esc.pivot, dict) else None
    if not guardado or not guardado.get("resultado"):
        return
    cfg = guardado.get("config") or {}
    hoja = _hoja(wb, "Pivot Registro de Cargas")
    fila = _portada(hoja, esc, "Registro de Cargas (Pivot)", _resumen_cfg(cfg))
    _tabla_pivot(hoja, fila, guardado["resultado"], cfg, "Registro de Cargas")
    _anchos(hoja, maximo=26)


def _resumen_cfg(cfg):
    def nombres(lista):
        return ", ".join(x["campo"] + (" ⚑" if x.get("incluir") else "") for x in lista) or "–"
    return (f"Filas: {nombres(cfg.get('filas', []))} · Columnas: {nombres(cfg.get('columnas', []))} · "
            f"Filtros: {nombres(cfg.get('filtros', []))}")


def _hoja_parametros(wb, esc):
    pa = _hoja(wb, "Parámetros")
    fila = _portada(pa, esc, "Parámetros del Escenario",
                    f"Año: {esc.anio or '–'} · Días del periodo: {esc.dias:g} · Réplica: {esc.replica}")
    _subtitulo(pa, "Archivos de Entrada", fila)
    fila += 1
    datos = [[a["nombre"], a.get("filas"), a.get("columnas"), "Sí" if a.get("incluido", True) else "No"]
             for a in esc.archivos]
    fila = _tabla(pa, fila, [("n", "Archivo", "texto"), ("f", "Filas", "entero"), ("c", "Columnas", "entero"),
                             ("i", "Incluido", "texto")], datos)
    for clase in modelo.ORDEN_CLASES:
        cfg = esc.clases.get(clase)
        if not cfg:
            continue
        info = modelo.CLASES[clase]
        _subtitulo(pa, f"Configuración de Tiempos – {info['plural']}", fila)
        fila += 1
        datos = [[col, modelo.NOMBRE_ESTADO.get(cfg["mapa_estados"].get(col), "Sin Asignar")]
                 for col in modelo.ordenar_columnas(clase, list(cfg["mapa_estados"]))]
        fila = _tabla(pa, fila, [("c", "Columna del Archivo", "texto"), ("e", "Estado", "texto")], datos)
        _subtitulo(pa, f"Configuración de Equipos – {info['plural']}", fila)
        fila += 1
        categoria = cfg.get("energia") or {}
        cols = [("f", "Flota", "texto"), ("n", f"N° de {info['plural']}", "entero"), ("t", info["tipo_titulo"], "texto")]
        if clase == "perforadoras":
            cols.append(("e", modelo.CATEGORIA, "texto"))
        datos = [[f, cfg["flotas"].get(f), cfg["tipos"].get(f) or modelo.SIN_ASIGNAR] +
                 ([categoria.get(f, "")] if clase == "perforadoras" else [])
                 for f in modelo.ordenar_flotas(clase, cfg["flotas"], cfg)]
        fila = _tabla(pa, fila, cols, datos)
        orden = cfg.get("orden") or {}
        if orden:
            nombres = dict(modelo.dims_orden(clase))
            _texto(pa, fila - 1, 1, "Orden de Reporte: " + "  ▸  ".join(
                f"{nombres.get(d, d)} ({', '.join(orden.get(d, []))})" for d in orden.get("prioridad", [])), suave=True)
            fila += 1
    if esc.destinos:
        _subtitulo(pa, "Tipo de Destino", fila)
        fila = _tabla(pa, fila + 1, [("d", "Destino", "texto"), ("t", "Tipo de Destino", "texto"),
                                     ("m", "Material (Mt)", "mt")],
                      [[d, c.get("tipo"), (c.get("carga") or 0) / 1e6] for d, c in esc.destinos.items()])
    if esc.origenes:
        _subtitulo(pa, "Tipo de Origen y Condición", fila)
        fila = _tabla(pa, fila + 1, [("o", "Origen", "texto"), ("f", "Fase", "texto"), ("t", "Tipo de Origen", "texto"),
                                     ("c", "Condición", "texto"), ("m", "Material (Mt)", "mt")],
                      [[o, c.get("fase"), c.get("tipo"), c.get("condicion"), (c.get("carga") or 0) / 1e6]
                       for o, c in esc.origenes.items()])
    _anchos(pa)


def exportar_pivot(ruta, esc, tabla, cfg=None):
    """Exporta el resultado de una tabla dinámica (con su gráfico)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tabla Dinámica"
    fila = _portada(ws, esc, "Tabla Dinámica", _resumen_cfg(cfg or {}))
    _tabla_pivot(ws, fila, tabla, cfg or {}, "Tabla Dinámica")
    _anchos(ws)
    ws.freeze_panes = ws.cell(row=fila + 1, column=tabla["n_dim"] + 1)
    guardar(wb, ruta)
    return ruta
