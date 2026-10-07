"""Exportación «Resultados [Escenario].xlsx»: Vector Plan vs Simulación por réplica.

Replica el formato del libro de referencia (colores, bordes, ajuste de texto) y se adapta a la
estructura del Vector Plan importado (secciones e ítems en el mismo orden y con los mismos textos).
Las variaciones son fórmulas de Excel funcionales: B/A, B-A o (B-A)/A según cada ítem.
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import plan

BLANCO = Side(style="medium", color="FFFFFF")
BORDE = Border(left=BLANCO, right=BLANCO, top=BLANCO, bottom=BLANCO)
F_SECCION = "004749"
F_PLAN = "464444"
F_ESC = "004A82"
F_VAR = "028411"
F_ITEM = "E7ECED"
F_CAB_PLAN = "3A3838"
F_CAB_ESC = "002060"
FUENTE_T = "Aptos Narrow"
FUENTE = "Calibri"
FILA0 = 9
COL_B, COL_C, COL_E, COL_G = 2, 3, 5, 7


def _relleno(c):
    return PatternFill("solid", start_color=c, end_color=c)


def _fecha_corrida(esc):
    """Fecha de la corrida del escenario (la de emisión si no se conoce)."""
    try:
        d = datetime.fromisoformat((esc.resultados or {}).get("fecha"))
    except (TypeError, ValueError):
        d = datetime.now()
    return d.strftime("%H:%M %d/%m/%Y")


def columnas_replica(j):
    """(columna valor, columna variación) de la réplica j (0, 1, …)."""
    base = COL_G + 4 * j
    return base, base + 2


def exportar(ruta, esc):
    if not esc.plan:
        raise ValueError("No hay un Vector Plan importado.")
    if not plan.replicas_resultados(esc):
        raise ValueError("No hay resultados para exportar.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    escribir(wb, ws, esc, "Gráficos")
    wb.save(ruta)
    return ruta


def escribir(wb, ws, esc, nombre_hoja_graficos):
    """Escribe la tabla Plan vs Simulación en 'ws' y sus gráficos (vinculados) en otra hoja."""
    reps = plan.replicas_resultados(esc)
    por = esc.resultados["por_replica"]
    nombre = esc.nombre_visible
    multi = len(reps) > 1
    ws.sheet_view.zoomScale = 85

    ws["B2"] = "Resultados de Simulación"
    ws["B2"].font = Font(name=FUENTE_T, size=16, bold=True, color="C00000")
    ws.row_dimensions[2].height = 21
    ws["B3"] = f"Caso {nombre}"
    ws["B3"].font = Font(name=FUENTE_T, size=11, bold=True, color="0000FF")
    ws["B4"] = f"Fecha de Corrida: {_fecha_corrida(esc)}"
    ws["B4"].font = Font(name=FUENTE_T, size=11)
    if multi:
        ws["B5"] = f"N° de Réplicas: {len(reps)}"
        ws["B5"].font = Font(name=FUENTE_T, size=11)

    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c = ws.cell(7, COL_E, esc.plan.get("nombre", "Plan"))
    c.font, c.fill, c.alignment = Font(name=FUENTE_T, size=11, color="FFFFFF"), _relleno(F_CAB_PLAN), centro
    for j, r in enumerate(reps):
        cv, cx = columnas_replica(j)
        c = ws.cell(7, cv, f"{nombre}\nRéplica {r}" if multi else nombre)
        c.font, c.fill, c.alignment = Font(name=FUENTE_T, size=11, bold=True, color="FFFFFF"), _relleno(F_CAB_ESC), centro
        c = ws.cell(7, cx, f"Var. {nombre}\nRéplica {r}" if multi else "Variación")
        c.font, c.alignment = Font(name=FUENTE_T, size=11, bold=True), centro
    ws.row_dimensions[7].height = 45 if multi or len(nombre) > 26 else 30

    fila = FILA0
    secciones = []
    for f in esc.plan["filas"]:
        ws.row_dimensions[fila].height = 24
        if f["tipo"] == "seccion":
            _seccion(ws, fila, f["texto"], len(reps))
            secciones.append((f["texto"], []))
        else:
            sims = [plan.valor(esc, f.get("vinculo"), por[r]) for r in reps]
            _item(ws, fila, f, sims)
            if secciones:
                secciones[-1][1].append((fila, f))
        fila += 1

    anchos = {"A": 2.5, "B": 45.86, "C": 8.43, "D": 7.86, "E": 16.43, "F": 7.86}
    for k, v in anchos.items():
        ws.column_dimensions[k].width = v
    for j in range(len(reps)):
        cv, cx = columnas_replica(j)
        ws.column_dimensions[get_column_letter(cv)].width = 22 if multi else 30
        ws.column_dimensions[get_column_letter(cv + 1)].width = 4.29
        ws.column_dimensions[get_column_letter(cx)].width = 14.57
        ws.column_dimensions[get_column_letter(cx + 1)].width = 4.29
    ws.freeze_panes = ws.cell(FILA0, COL_C)
    _graficos(wb, ws, secciones, reps, multi, nombre, esc.plan.get("nombre", "Plan"), nombre_hoja_graficos)


def _seccion(ws, fila, texto, n_reps):
    c = ws.cell(fila, COL_B, texto)
    c.font = Font(name=FUENTE, size=9, bold=True, color="FFFFFF")
    c.fill = _relleno(F_SECCION)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    c.border = BORDE
    rellenos = [(COL_C, F_SECCION), (COL_E, F_PLAN)]
    for j in range(n_reps):
        cv, cx = columnas_replica(j)
        rellenos += [(cv, F_ESC), (cx, F_VAR)]
    for col, color in rellenos:
        c = ws.cell(fila, col)
        c.fill, c.border = _relleno(color), BORDE


def _item(ws, fila, f, sims):
    fuente = Font(name=FUENTE, size=9, color="000000")
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c = ws.cell(fila, COL_B, f["texto"])
    c.font, c.fill, c.border = fuente, _relleno(F_ITEM), BORDE
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=2)
    c = ws.cell(fila, COL_C, f["unidad"])
    c.font, c.fill, c.border, c.alignment = fuente, _relleno(F_ITEM), BORDE, centro
    fmt_v = plan.formato_excel(f)
    fmt_x = plan.formato_excel(f, "var")
    e = ws.cell(fila, COL_E, f["valor"] if f["valor"] is not None else "-")
    e.font, e.fill, e.border, e.alignment = fuente, _relleno(F_ITEM), BORDE, centro
    e.number_format = fmt_v
    le = get_column_letter(COL_E)
    for j, s in enumerate(sims):
        cv, cx = columnas_replica(j)
        g = ws.cell(fila, cv, round(s, 6) if s is not None else "-")
        g.font, g.fill, g.border, g.alignment = fuente, _relleno(F_ITEM), BORDE, centro
        g.number_format = fmt_v
        lg = get_column_letter(cv)
        a, b = f"{le}{fila}", f"{lg}{fila}"
        tipo = f.get("variacion") or "resta"
        formula = {"div": f'=IFERROR({b}/{a},"-")', "dif": f'=IFERROR(({b}-{a})/{a},"-")'}.get(
            tipo, f'=IFERROR({b}-{a},"-")')
        x = ws.cell(fila, cx, formula)
        x.font, x.fill, x.border, x.alignment = fuente, _relleno(F_ITEM), BORDE, centro
        x.number_format = fmt_x


def _graficos(wb, ws, secciones, reps, multi, nombre, nombre_plan, nombre_hoja="Gráficos"):
    hg = wb.create_sheet(nombre_hoja)
    hg.sheet_view.showGridLines = False
    hg["B2"] = "Gráficos Comparativos – Plan vs Simulación"
    hg["B2"].font = Font(name=FUENTE_T, size=14, bold=True, color="C00000")
    hg["B3"] = f"Caso {nombre}"
    hg["B3"].font = Font(name=FUENTE_T, size=11, bold=True, color="0000FF")
    hg["B4"] = ws["B4"].value
    hg["B4"].font = Font(name=FUENTE_T, size=11)
    colores = ["004A82", "028411", "BF8700", "8250DF", "BC4C00", "1B7C83"]
    graficos = []
    for titulo, items in secciones:
        if not items:
            continue
        g = BarChart()
        g.type = "bar"
        g.grouping = "clustered"
        g.title = titulo.strip()
        g.title.overlay = False
        g.height, g.width = 6.5 + 0.35 * len(items) * (1 + len(reps)) ** 0.5, 15.5
        g.y_axis.majorGridlines = None
        g.x_axis.delete = False
        g.y_axis.delete = False
        g.x_axis.scaling.orientation = "maxMin"          # mismo orden que la tabla
        g.y_axis.numFmt = plan.formato_excel(items[0][1])
        g.legend.position = "b"
        g.legend.overlay = False
        f0, f1 = items[0][0], items[-1][0]
        cats = Reference(ws, min_col=COL_B, min_row=f0, max_row=f1)
        series = [(COL_E, nombre_plan, "6E7781")]
        for j, r in enumerate(reps):
            series.append((columnas_replica(j)[0], f"Réplica {r}" if multi else nombre, colores[j % len(colores)]))
        from openpyxl.chart import Series
        for col, etq, color in series:
            s = Series(Reference(ws, min_col=col, min_row=f0, max_row=f1), title=etq)
            s.graphicalProperties.solidFill = color
            s.graphicalProperties.line.solidFill = color
            g.series.append(s)
        g.set_categories(cats)
        graficos.append(g)
    # Dos gráficos por fila, sin superponerse (alto de fila ≈ 0.53 cm)
    fila = 6
    for k in range(0, len(graficos), 2):
        par = graficos[k:k + 2]
        for i, g in enumerate(par):
            hg.add_chart(g, f"{'B' if i == 0 else 'K'}{fila}")
        fila += int(max(g.height for g in par) / 0.53) + 3
