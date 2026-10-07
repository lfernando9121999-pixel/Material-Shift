"""Gráficos nativos (Canvas): barras agrupadas, apiladas o en cascada (verticales u horizontales) y gráfico
de rango (puntos Mínimo · Promedio · Máximo unidos por una línea).

Se adaptan al modo claro/oscuro, se redibujan una sola vez por cambio de tamaño (sin parpadeo), rotulan
los valores dentro de cada barra cuando hay espacio y muestran el valor al pasar el mouse.
Clic derecho: Copiar Datos de Gráfico · Copiar Gráfico como Imagen · Copiar Gráfico como Formato de PPT
(gráfico nativo y editable de PowerPoint).
"""
import math
import tkinter as tk

from tema import T

PALETA_SERIES_CLARO = ["#0969DA", "#1A7F37", "#BF8700", "#8250DF", "#BC4C00", "#1B7C83",
                       "#CF222E", "#6E7781", "#4D2D00", "#0550AE", "#116329", "#953800"]
PALETA_SERIES_OSCURO = ["#4493F8", "#3FB950", "#D29922", "#A371F7", "#DB6D28", "#39C5CF",
                        "#F85149", "#8B949E", "#E3B341", "#79C0FF", "#56D364", "#FFA657"]


def paleta_series():
    return PALETA_SERIES_OSCURO if T.nombre == "oscuro" else PALETA_SERIES_CLARO


def _paso_bonito(maximo, n=5):
    if maximo <= 0:
        return 1.0
    crudo = maximo / n
    mag = 10 ** math.floor(math.log10(crudo))
    for m in (1, 2, 2.5, 5, 10):
        if crudo <= m * mag:
            return m * mag
    return 10 * mag


def _texto_sobre(color):
    """Color de texto legible sobre un relleno (blanco u oscuro según la luminancia)."""
    try:
        c = color.lstrip("#")
        r, g, b = (int(c[i:i + 2], 16) / 255 for i in (0, 2, 4))
        lum = 0.2126 * r + 0.7152 * g + 0.0722 * b
        return "#0D1117" if lum > 0.55 else "#FFFFFF"
    except (ValueError, AttributeError):
        return "#FFFFFF"


def copiar_imagen(img):
    """Compatibilidad: copia una imagen PIL al portapapeles."""
    import portapapeles
    portapapeles.copiar_imagen(img)


class GraficoBarras(tk.Canvas):
    def __init__(self, master, alto=320):
        super().__init__(master, height=T.px(alto), highlightthickness=0, borderwidth=0)
        self.categorias, self.series = [], []
        self.apilado, self.horizontal, self.cascada, self.rango = False, False, False, False
        self.titulo, self.unidad = "", ""
        self.nombre = ""                   # título del gráfico al copiarlo a PowerPoint
        self.formato = lambda v: f"{v:,.1f}"
        self.formato_eje = None
        self.colores_barra = None
        self.mostrar_valores = True
        self.paso, self.tope = None, None
        self.mostrar_total = True
        self.notas = None
        self._tip = None
        self._id = None
        self._tam = None
        self.bind("<Configure>", self._al_configurar)
        self.bind("<Button-3>", self._menu)
        T.registrar(self.dibujar)

    def _al_configurar(self, e):
        tam = (e.width, e.height)
        if tam == self._tam:
            return                          # la ventana solo se movió: no hay nada que redibujar
        if self._tam is None or self._tam[0] < 50:
            self._programar()
        else:
            self._programar(70)             # cambio de tamaño: primero se pinta la ventana, luego el gráfico

    def _programar(self, retardo=0):
        if self._id is None:
            self._id = self.after(retardo, self.dibujar) if retardo else self.after_idle(self.dibujar)

    def datos(self, categorias, series, apilado=False, horizontal=False, titulo="", unidad="",
              formato=None, colores_barra=None, mostrar_valores=True, paso=None, tope=None,
              formato_eje=None, cascada=False, mostrar_total=True, notas=None, rango=False):
        """series: lista de (nombre, valores, color|None). Una categoría puede ser (nombre, sufijo).
        rango=True: series = Mínimo, Promedio y Máximo (gráfico de puntos por categoría, horizontal)."""
        self.notas = list(notas) if notas else None
        self.categorias = [c if isinstance(c, tuple) else str(c) for c in categorias]
        paleta = paleta_series()
        self.series = [(n, [0.0 if v is None else float(v) for v in vals], c or paleta[i % len(paleta)])
                       for i, (n, vals, c) in enumerate(series)]
        self.apilado, self.horizontal, self.cascada, self.rango = apilado, horizontal or rango, cascada, rango
        self.titulo, self.unidad = titulo, unidad
        if formato:
            self.formato = formato
        self.formato_eje = formato_eje
        self.colores_barra = colores_barra
        self.mostrar_valores = mostrar_valores
        self.paso, self.tope = paso, tope
        self.mostrar_total = mostrar_total
        self.dibujar()

    # ------------------------------------------------------------------
    def _nombre(self, cat):
        return cat[0] if isinstance(cat, tuple) else cat

    def dibujar(self):
        self._id = None
        p = T.p
        try:
            self.configure(background=p["panel"])
        except tk.TclError:
            return
        self.delete("all")
        self._tip = None
        W, H = self.winfo_width(), self.winfo_height()
        self._tam = (W, H)
        if W < 50 or not self.categorias or not self.series:
            if W > 50 and H > 20:
                self.create_text(W / 2, H / 2, text="Sin datos para graficar", fill=p["suave"], font=T.f["base"])
            return
        f_chica = T.f["chica"]
        px = T.px
        top = px(8)
        if self.titulo:
            self.create_text(px(10), top, text=self.titulo, anchor="nw", fill=p["texto"], font=T.f["negrita"])
            top += px(20)
        if (len(self.series) > 1 or self.colores_barra is None) and not self.cascada and not self.rango:
            x = px(10)
            for nombre, _v, color in self.series:
                ancho = f_chica.measure(nombre) + px(24)
                if x + ancho > W - px(10):
                    x = px(10)
                    top += px(16)
                self.create_rectangle(x, top + px(3), x + px(10), top + px(13), fill=color, outline="")
                self.create_text(x + px(14), top + px(8), text=nombre, anchor="w", fill=p["suave"], font=f_chica)
                x += ancho
            top += px(20)
        n_cat = len(self.categorias)
        if self.rango:
            maximos = [max(s[1][i] for s in self.series) for i in range(n_cat)]
        elif self.cascada:
            vals = self.series[0][1]
            acumulado, maximos = 0.0, []
            for i, v in enumerate(vals):
                if i == n_cat - 1:
                    maximos.append(v)
                else:
                    acumulado += max(0.0, v)
                    maximos.append(acumulado)
        elif self.apilado:
            maximos = [sum(max(0.0, s[1][i]) for s in self.series) for i in range(n_cat)]
        else:
            maximos = [max((max(0.0, s[1][i]) for s in self.series), default=0) for i in range(n_cat)]
        vmax = max(maximos + [0.0]) or 1.0
        paso = self.paso or _paso_bonito(vmax)
        tope = self.tope if self.tope and self.tope >= vmax * 0.999 else math.ceil(vmax / paso - 1e-6) * paso
        tope = max(tope, paso)
        f_eje = self.formato_eje or self.formato
        etq_val = [f_eje(paso * k) for k in range(int(round(tope / paso)) + 1)]

        if self.horizontal:
            nombres = [self._nombre(c) for c in self.categorias]
            sufijos = [c[1] if isinstance(c, tuple) else "" for c in self.categorias]
            ancho_suf = max((f_chica.measure(s) for s in sufijos), default=0)
            ancho_nom = max(f_chica.measure(n) for n in nombres)
            hueco = px(12) if ancho_suf else 0
            ancho_cat = min(ancho_nom + ancho_suf + hueco + px(14), int(W * 0.34))
            if self.rango:
                margen = 70
            else:
                margen = 52 if self.mostrar_valores and (self.mostrar_total or not self.apilado) else 18
            ancho_notas = max((f_chica.measure(t) for t in self.notas), default=0) + px(14) if self.notas else 0
            x0, x1 = px(10) + ancho_cat, W - px(margen) - ancho_notas
            y0, y1 = top + px(2), H - px(24)
            largo = x1 - x0
            for k, et in enumerate(etq_val):
                x = x0 + largo * (k * paso) / tope
                self.create_line(x, y0, x, y1, fill=p["rejilla"])
                self.create_text(x, y1 + px(5), text=et, anchor="n", fill=p["suave"], font=f_chica)
            banda = (y1 - y0) / n_cat
            for i in range(n_cat):
                yc = y0 + banda * (i + 0.5)
                if ancho_suf:
                    self.create_text(x0 - px(8), yc, text=sufijos[i], anchor="e", fill=p["suave"], font=f_chica)
                    self.create_text(x0 - px(8) - ancho_suf - hueco, yc,
                                     text=self._corta(nombres[i], ancho_nom + px(2), f_chica),
                                     anchor="e", fill=p["texto"], font=f_chica)
                else:
                    self.create_text(x0 - px(8), yc, text=self._corta(nombres[i], ancho_cat - px(10), f_chica),
                                     anchor="e", fill=p["texto"], font=f_chica)
                if self.rango:
                    self._rango_categoria(i, yc, banda, x0, largo, tope, x1)
                else:
                    self._barras_categoria(i, yc, banda, x0, largo, tope)
                if self.notas and i < len(self.notas):
                    self.create_text(W - px(8) - ancho_notas + px(14), yc, text=self.notas[i], anchor="w",
                                     fill=p["suave"], font=f_chica)
            self.create_line(x0, y0, x0, y1, fill=p["borde"])
        else:
            ancho_eje = max(f_chica.measure(e) for e in etq_val) + px(14)
            x0, x1 = px(8) + ancho_eje, W - px(10)
            etiquetas = [self._nombre(c) + (f" {c[1]}" if isinstance(c, tuple) else "") for c in self.categorias]
            ancho_etq = max(f_chica.measure(c) for c in etiquetas)
            salto = 1
            if n_cat > 24 and ancho_etq + px(8) > (x1 - x0) / n_cat:
                salto = math.ceil((ancho_etq + px(8)) / ((x1 - x0) / n_cat))
            rotar = salto == 1 and n_cat > 6 and ancho_etq > (x1 - x0) / n_cat - px(6)
            alto_etq = px(60) if rotar else px(22)
            y0, y1 = top + px(12), H - alto_etq
            alto = y1 - y0
            for k, et in enumerate(etq_val):
                y = y1 - alto * (k * paso) / tope
                self.create_line(x0, y, x1, y, fill=p["rejilla"])
                self.create_text(x0 - px(6), y, text=et, anchor="e", fill=p["suave"], font=f_chica)
            banda = (x1 - x0) / n_cat
            for i, cat in enumerate(etiquetas):
                xc = x0 + banda * (i + 0.5)
                if salto > 1:
                    if i % salto == 0:
                        self.create_text(xc, y1 + px(5), text=cat, anchor="n", fill=p["texto"], font=f_chica)
                elif rotar:
                    self.create_text(xc, y1 + px(5), text=self._corta(cat, alto_etq - px(8), f_chica),
                                     anchor="ne", angle=35, fill=p["texto"], font=f_chica)
                else:
                    self.create_text(xc, y1 + px(5), text=self._corta(cat, banda - px(4), f_chica),
                                     anchor="n", fill=p["texto"], font=f_chica)
                if self.cascada:
                    self._barra_cascada(i, xc, banda, y1, alto, tope)
                else:
                    self._barras_categoria(i, xc, banda, y1, alto, tope)
            self.create_line(x0, y1, x1, y1, fill=p["borde"])
        if self.unidad:
            if self.horizontal and self.notas:
                self.create_text(x1, px(8), text=self.unidad, anchor="ne", fill=p["suave"], font=f_chica)
            else:
                self.create_text(W - px(10), px(8), text=self.unidad, anchor="ne", fill=p["suave"], font=f_chica)

    # ------------------------------------------------------------------
    def _rango_categoria(self, i, yc, banda, base, largo, tope, limite):
        """Puntos Mínimo · Promedio · Máximo unidos por una línea discontinua, con sus valores sin
        superponerse (arriba el promedio; abajo el mínimo y el máximo, o al costado si no caben)."""
        p = T.p
        px = T.px
        f = T.f["chica"]
        fb = T.f["negrita"]
        mn, pr, mx = (s[1][i] for s in self.series[:3])
        xs = [base + largo * v / tope for v in (mn, pr, mx)]
        punto = p["texto"] if T.nombre == "oscuro" else "#12355B"
        linea = "#58A6FF" if T.nombre == "oscuro" else "#3B8ED8"
        cat = self._nombre(self.categorias[i])
        self.create_line(xs[0], yc, xs[2], yc, fill=linea, dash=(4, 3), width=2)
        r = min(px(6), max(px(3), banda * 0.16))
        nombres = ("Mínimo", "Promedio", "Máximo")
        for x, v, n in zip(xs, (mn, pr, mx), nombres):
            item = self.create_oval(x - r, yc - r, x + r, yc + r, fill=punto, outline=p["panel"], width=1)
            tip = f"{cat} · {n}: {self.formato(v)} {self.unidad}".strip()
            self.tag_bind(item, "<Enter>", lambda e=None, t=tip: self._mostrar_tip(e, t))
            self.tag_bind(item, "<Leave>", lambda *_: self._ocultar_tip())
        if not self.mostrar_valores:
            return
        t_mn, t_pr, t_mx = f"Mín. {self.formato(mn)}", f"Prom. {self.formato(pr)}", f"Máx. {self.formato(mx)}"
        alto_txt = f.metrics("linespace")
        arriba, abajo = yc - r - px(1), yc + r + px(1)
        dos_niveles = banda >= 2 * alto_txt + 2 * r + px(4)
        # promedio arriba (centrado en su punto, dentro del área)
        w_pr = fb.measure(t_pr)
        xp = min(max(xs[1], base + w_pr / 2), limite + px(60) - w_pr / 2)
        if dos_niveles:
            self.create_text(xp, arriba, text=t_pr, anchor="s", fill=p["texto"], font=fb)
            w_mn, w_mx = f.measure(t_mn), f.measure(t_mx)
            xa = max(xs[0], base + w_mn / 2)
            xb = xs[2]
            if xa + w_mn / 2 + px(6) > xb - w_mx / 2:          # se superpondrían: se separan hacia los lados
                xa = min(xa, xb - w_mx / 2 - px(6) - w_mn / 2)
                if xa - w_mn / 2 < base:
                    xa = base + w_mn / 2
                    xb = xa + w_mn / 2 + px(6) + w_mx / 2
            self.create_text(xa, abajo, text=t_mn, anchor="n", fill=p["suave"], font=f)
            self.create_text(xb, abajo, text=t_mx, anchor="n", fill=p["suave"], font=f)
        else:                                   # una sola línea: valores a la derecha del máximo
            self.create_text(xs[2] + r + px(6), yc, text=f"{t_mn}  ·  {t_pr}  ·  {t_mx}", anchor="w",
                             fill=p["suave"], font=f)

    def _barra_cascada(self, i, centro, banda, base, largo, tope):
        vals = self.series[0][1]
        ultimo = len(vals) - 1
        inicio = 0.0 if i == ultimo else sum(max(0.0, v) for v in vals[:i])
        v = max(0.0, vals[i])
        grosor = min(banda * 0.62, T.px(96))
        color = self.colores_barra[i] if self.colores_barra else self.series[0][2]
        nombre = self._nombre(self.categorias[i])
        self._rect(centro - grosor / 2, centro + grosor / 2, base, largo, tope, inicio, inicio + v, color,
                   f"{nombre}: {self.formato(v)} {self.unidad}".strip(), self.formato(v) if self.mostrar_valores else "",
                   afuera=True, centrar=True)
        if i < ultimo:
            y = base - largo * (inicio + v) / tope
            self.create_line(centro + grosor / 2, y, centro + banda - grosor / 2, y, fill=T.p["suave"], dash=(3, 3))

    def _barras_categoria(self, i, centro, banda, base, largo, tope):
        px = T.px
        ns = len(self.series)
        ocupado = banda * 0.74
        cat = self._nombre(self.categorias[i])
        if self.apilado or ns == 1:
            grosor = min(ocupado, px(34 if self.horizontal else 64))     # barras proporcionadas aunque sobre alto
            acumulado = 0.0
            for j, (nombre, vals, color) in enumerate(self.series):
                v = max(0.0, vals[i])
                if v <= 0:
                    continue
                c = self.colores_barra[i] if (self.colores_barra and ns == 1) else color
                a, b = acumulado, acumulado + v
                acumulado = b if self.apilado else 0.0
                etq = self.formato(v) if self.mostrar_valores else ""
                self._rect(centro - grosor / 2, centro + grosor / 2, base, largo, tope, a, b, c,
                           f"{cat} · {nombre}: {self.formato(v)} {self.unidad}".strip(), etq,
                           afuera=(ns == 1))
            if self.mostrar_valores and self.mostrar_total and self.apilado and ns > 1 and acumulado > 0:
                self._etiqueta_total(centro, base, largo, tope, acumulado, grosor)
        else:
            grosor = min(ocupado / ns, px(40))
            inicio = centro - grosor * ns / 2
            for j, (nombre, vals, color) in enumerate(self.series):
                v = max(0.0, vals[i])
                c0 = inicio + j * grosor
                etq = self.formato(v) if self.mostrar_valores else ""
                self._rect(c0 + 1, c0 + grosor - 1, base, largo, tope, 0.0, v, color,
                           f"{cat} · {nombre}: {self.formato(v)} {self.unidad}".strip(), etq, afuera=True)

    def _rect(self, c0, c1, base, largo, tope, a, b, color, tip, etiqueta="", afuera=False, centrar=False):
        f = T.f["chica"]
        alto_txt = f.metrics("linespace")
        if self.horizontal:
            xa, xb = base + largo * a / tope, base + largo * b / tope
            item = self.create_rectangle(xa, c0, xb, c1, fill=color, outline=T.p["panel"])
            if etiqueta:
                w = f.measure(etiqueta)
                if w + T.px(8) <= xb - xa and alto_txt <= c1 - c0 + 2:
                    self.create_text((xa + xb) / 2 if not afuera else xb - T.px(5), (c0 + c1) / 2, text=etiqueta,
                                     anchor="center" if not afuera else "e", fill=_texto_sobre(color), font=f)
                elif afuera:
                    self.create_text(xb + T.px(4), (c0 + c1) / 2, text=etiqueta, anchor="w",
                                     fill=T.p["texto"], font=f)
        else:
            ya, yb = base - largo * a / tope, base - largo * b / tope
            item = self.create_rectangle(c0, yb, c1, ya, fill=color, outline=T.p["panel"])
            if etiqueta:
                w = f.measure(etiqueta)
                cabe = w + T.px(4) <= c1 - c0 and alto_txt + T.px(4) <= ya - yb
                if cabe and (centrar or not afuera):
                    self.create_text((c0 + c1) / 2, (ya + yb) / 2, text=etiqueta, anchor="center",
                                     fill=_texto_sobre(color), font=f)
                elif cabe:
                    self.create_text((c0 + c1) / 2, yb + T.px(3), text=etiqueta, anchor="n",
                                     fill=_texto_sobre(color), font=f)
                elif afuera and w <= (c1 - c0) + T.px(16):
                    self.create_text((c0 + c1) / 2, yb - T.px(2), text=etiqueta, anchor="s",
                                     fill=T.p["texto"], font=f)
        self.tag_bind(item, "<Enter>", lambda e=None, t=tip: self._mostrar_tip(e, t))
        self.tag_bind(item, "<Motion>", lambda e=None, t=tip: self._mostrar_tip(e, t))
        self.tag_bind(item, "<Leave>", lambda *_: self._ocultar_tip())

    def _etiqueta_total(self, centro, base, largo, tope, valor, grosor):
        p = T.p
        texto = self.formato(valor)
        if self.horizontal:
            x = base + largo * valor / tope + T.px(4)
            self.create_text(x, centro, text=texto, anchor="w", fill=p["texto"], font=T.f["negrita"])
        else:
            y = base - largo * valor / tope - T.px(3)
            if T.f["chica"].measure(texto) <= grosor + T.px(18):
                self.create_text(centro, y, text=texto, anchor="s", fill=p["texto"], font=T.f["negrita"])

    @staticmethod
    def _corta(texto, ancho, fuente):
        if fuente.measure(texto) <= ancho:
            return texto
        while texto and fuente.measure(texto + "…") > ancho:
            texto = texto[:-1]
        return texto + "…"

    def _mostrar_tip(self, e, texto):
        p = T.p
        self._ocultar_tip()
        if e is None or not hasattr(e, "x"):       # algunos eventos sintéticos de Tk llegan sin datos
            try:
                e = type("E", (), {"x": self.winfo_pointerx() - self.winfo_rootx(),
                                   "y": self.winfo_pointery() - self.winfo_rooty()})()
            except tk.TclError:
                return
        x, y = e.x + T.px(12), e.y - T.px(10)
        t = self.create_text(x + T.px(8), y, text=texto, anchor="w", fill=p["texto"], font=T.f["chica"], tags=("tip",))
        x0, y0, x1, y1 = self.bbox(t)
        if x1 > self.winfo_width() - 4:
            self.move(t, -(x1 - x0) - T.px(36), 0)
            x0, y0, x1, y1 = self.bbox(t)
        r = self.create_rectangle(x0 - T.px(6), y0 - T.px(4), x1 + T.px(6), y1 + T.px(4),
                                  fill=p["panel"], outline=p["borde"], tags=("tip",))
        self.tag_raise(t, r)
        self._tip = (t, r)

    def _ocultar_tip(self):
        if self._tip:
            for i in self._tip:
                self.delete(i)
            self._tip = None

    # ---- copiar --------------------------------------------------------------------
    def _menu(self, e):
        if not self.categorias or not self.series:
            return
        from ui_comun import mostrar_menu, nuevo_menu
        m = nuevo_menu(self)
        m.add_command(label="Copiar Datos de Gráfico", command=self.copiar_datos)
        m.add_command(label="Copiar Gráfico como Imagen", command=self.copiar_grafico)
        m.add_command(label="Copiar Gráfico como Formato de PPT", command=self.copiar_ppt)
        mostrar_menu(m, e)

    def _etiquetas_categoria(self):
        return [self._nombre(c) + (f" – {c[1]}" if isinstance(c, tuple) and c[1] else "") for c in self.categorias]

    def copiar_datos(self):
        import portapapeles
        from ui_comun import avisar_estado
        cab = ["Categoría"] + [n for n, _v, _c in self.series]
        filas = [[c] + [f"{s[1][i]:.6g}" for s in self.series] for i, c in enumerate(self._etiquetas_categoria())]
        portapapeles.copiar_texto(portapapeles.tsv(cab, filas))
        avisar_estado(self, "Datos del gráfico copiados (péguelos en Excel)")

    def copiar_grafico(self):
        import portapapeles
        import render
        from ui_comun import avisar_estado
        self._ocultar_tip()
        try:
            portapapeles.copiar_imagen(render.lienzo_a_imagen(self))
            avisar_estado(self, "Gráfico copiado como imagen")
        except Exception as ex:  # noqa
            avisar_estado(self, f"No se pudo copiar el gráfico: {ex}")

    def copiar_ppt(self):
        import portapapeles
        from ui_comun import avisar_estado
        try:
            ruta = portapapeles.ruta_temporal(".pptx")
            self.a_pptx(ruta)
        except Exception as ex:  # noqa
            avisar_estado(self, f"No se pudo preparar el gráfico: {ex}")
            return
        avisar_estado(self, "Preparando el gráfico para PowerPoint…")

        def listo(ok, msg):
            texto = ("Gráfico copiado en formato de PowerPoint: péguelo con Ctrl+V (es editable)" if ok else
                     "No se pudo copiar en formato de PowerPoint (¿está instalado PowerPoint?). Use «Copiar Gráfico "
                     "como Imagen».")
            try:
                self.after(0, lambda: avisar_estado(self, texto))
            except (tk.TclError, RuntimeError):
                pass
        portapapeles.copiar_pptx_en_segundo_plano(ruta, listo)

    def a_pptx(self, ruta):
        """Lámina con el gráfico nativo de PowerPoint (mismos datos, colores y orientación)."""
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.util import Inches, Pt
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        cats = self._etiquetas_categoria()
        cd = CategoryChartData()
        cd.categories = cats
        if self.cascada:
            vals = self.series[0][1]
            base, acum = [], 0.0
            for i, v in enumerate(vals):
                base.append(0.0 if i == len(vals) - 1 else acum)
                acum += v if i < len(vals) - 1 else 0
            cd.add_series("Base", base)
            cd.add_series(self.series[0][0], vals)
            tipo = XL_CHART_TYPE.COLUMN_STACKED
        elif self.rango:
            for n, v, _c in self.series[:3]:
                cd.add_series(n, v)
            tipo = XL_CHART_TYPE.LINE_MARKERS
        else:
            for n, v, _c in self.series:
                cd.add_series(n, v)
            if self.horizontal:
                tipo = XL_CHART_TYPE.BAR_STACKED if self.apilado else XL_CHART_TYPE.BAR_CLUSTERED
            else:
                tipo = XL_CHART_TYPE.COLUMN_STACKED if self.apilado else XL_CHART_TYPE.COLUMN_CLUSTERED
        w = max(1, self.winfo_width())
        h = max(1, self.winfo_height())
        ancho = 11.5
        alto = max(2.6, min(6.6, ancho * h / w))
        gf = s.shapes.add_chart(tipo, Inches(0.9), Inches(0.45), Inches(ancho), Inches(alto), cd)
        ch = gf.chart
        ch.font.size = Pt(10)
        ch.font.name = "Segoe UI"
        if self.nombre:
            ch.has_title = True
            ch.chart_title.text_frame.text = self.nombre
            ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
        else:
            ch.has_title = False
        n_series = len(cd._series) if hasattr(cd, "_series") else len(self.series)
        ch.has_legend = (n_series > 1 and not self.cascada) or self.rango
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        if self.horizontal and not self.rango:
            ch.category_axis.reverse_order = True          # mismo orden que en pantalla
        ch.value_axis.has_major_gridlines = True
        ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xE1, 0xE4, 0xE8)
        muestra = [v for _n, vals, _c in self.series for v in vals if v]
        fmt_num = "#,##0" if muestra and max(abs(v) for v in muestra) >= 100 else "#,##0.00"
        ch.value_axis.tick_labels.number_format = fmt_num
        ch.value_axis.tick_labels.number_format_is_linked = False

        def rgb(c):
            c = c.lstrip("#")
            return RGBColor(int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16))
        series_ppt = list(ch.series)
        if self.cascada:
            series_ppt[0].format.fill.background()
            serie = series_ppt[1]
            for i in range(len(cats)):
                pt = serie.points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(self.colores_barra[i] if self.colores_barra else self.series[0][2])
            serie.data_labels.show_value = True
            serie.data_labels.number_format = fmt_num
            serie.data_labels.number_format_is_linked = False
        elif self.rango:
            for sp in series_ppt:
                sp.format.line.fill.background()
                sp.marker.size = 9
                sp.smooth = False
                sp.data_labels.show_value = True
                sp.data_labels.number_format = "#,##0.00"
                sp.data_labels.number_format_is_linked = False
        else:
            for sp, (_n, _v, color) in zip(series_ppt, self.series):
                sp.format.fill.solid()
                sp.format.fill.fore_color.rgb = rgb(color)
            if len(self.series) == 1 and self.colores_barra:
                for i, c in enumerate(self.colores_barra[:len(cats)]):
                    pt = series_ppt[0].points[i]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = rgb(c)
            if self.mostrar_valores and len(cats) <= 40:
                for sp in series_ppt:
                    sp.data_labels.show_value = True
                    sp.data_labels.font.size = Pt(9)
                    sp.data_labels.number_format = fmt_num
                    sp.data_labels.number_format_is_linked = False
        prs.save(ruta)
        return ruta
