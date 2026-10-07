"""Exportación del reporte PowerPoint (16:9) con la configuración base (sin filtros).

Una lámina por pestaña o subpestaña cuando el contenido lo permite. Las tablas largas se resumen o se
agrupan (nunca se dividen en varias láminas) y los gráficos son nativos de PowerPoint (editables, con
sus datos en Excel).
"""
from datetime import datetime

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import config
import fmt
import modelo
import reporte as rp

FUENTE = "Segoe UI"
ANCHO, ALTO = 13.333, 7.5
SERIES = ["0969DA", "1A7F37", "BF8700", "8250DF", "BC4C00", "1B7C83", "CF222E", "6E7781"]
COLOR_TIPO = {"Expit": rp.COLOR_ESTADO["TP"], "Rehandle": rp.COLOR_ESTADO["DPP"], "N/A": rp.COLOR_ESTADO["SB"],
              "Chancador": rp.COLOR_ESTADO["DENP"], "Stockpile": rp.COLOR_ESTADO["DEP"],
              "Botadero": rp.COLOR_ESTADO["DPNP"], "Total": "57606A"}
GRUPOS = ["EEF4FC", "EDF7F0", "FCF7EA", "F3F0FC", "ECF6F7", "FCF1F0"]
MAX_FILAS = 26          # filas de tabla que caben legibles en una lámina; sobre eso se resume


def _rgb(hex_):
    return RGBColor.from_string(hex_.upper())


def _formato(tipo, v):
    if v is None or v == "":
        return "" if tipo == "texto" else fmt.NODATO
    if tipo == "texto" or isinstance(v, str):
        return str(v)
    if tipo == "entero":
        return fmt.entero(v)
    if tipo == "dec1":
        return fmt.decimal1(v)
    if tipo == "pct":
        return fmt.porcentaje(v)
    if tipo == "dec2":
        return f"{float(v):,.2f}"
    return str(v)


def _texto(slide, x, y, w, h, texto, tam=12, color=rp.COLOR_TEXTO, negrita=False, alinear=PP_ALIGN.LEFT):
    caja = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = caja.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = alinear
    r = p.add_run()
    r.text = texto
    r.font.name, r.font.size, r.font.bold = FUENTE, Pt(tam), negrita
    r.font.color.rgb = _rgb(color)
    return caja


def _rect(slide, x, y, w, h, color, forma=MSO_SHAPE.RECTANGLE, borde=None):
    f = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    f.fill.solid()
    f.fill.fore_color.rgb = _rgb(color)
    if borde:
        f.line.color.rgb = _rgb(borde)
        f.line.width = Pt(0.75)
    else:
        f.line.fill.background()
    f.shadow.inherit = False
    return f


def fecha_corrida(esc):
    try:
        return datetime.fromisoformat((esc.resultados or {}).get("fecha")).strftime("%H:%M %d/%m/%Y")
    except (TypeError, ValueError):
        return datetime.now().strftime("%H:%M %d/%m/%Y")


class Reporte:
    def __init__(self, esc):
        self.esc = esc
        self.res = esc.resultados
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(ANCHO), Inches(ALTO)
        self.vacia = self.prs.slide_layouts[6]
        self.fecha = fecha_corrida(esc)

    # ---- estructura -------------------------------------------------------------------------------------------
    def portada(self):
        s = self.prs.slides.add_slide(self.vacia)
        _rect(s, 0, 0, ANCHO, ALTO, rp.COLOR_CAB)
        _rect(s, 0.8, 3.15, 1.2, 0.06, "9DB8D2")
        try:
            s.shapes.add_picture(str(config.dir_recursos() / "icono.png"), Inches(0.75), Inches(0.55),
                                 Inches(1.0), Inches(1.0))
        except Exception:  # noqa
            pass
        _texto(s, 0.7, 1.75, 11.5, 0.9, config.NOMBRE_PROGRAMA, 38, "FFFFFF", True)
        _texto(s, 0.7, 2.55, 11.5, 0.5, "Reporte de Resultados de Simulación", 20, "C9D6E3")
        _texto(s, 0.7, 3.45, 11.8, 0.7, f"Caso {self.esc.nombre_visible}", 26, "FFFFFF", True)
        _texto(s, 0.7, 4.1, 11.8, 0.45, f"Fecha de Corrida: {self.fecha}", 15, "C9D6E3")
        if self.esc.descripcion.strip():
            _texto(s, 0.7, 4.75, 11.8, 1.2, self.esc.descripcion.strip()[:400], 12, "C9D6E3")
        _texto(s, 0.7, 6.55, 11.8, 0.4, f"{config.NOMBRE_PROGRAMA} {config.VERSION_CORTA}  ·  Emitido el {fmt.fecha()}",
               11, "9DB8D2")

    def lamina(self, titulo, subtitulo=""):
        s = self.prs.slides.add_slide(self.vacia)
        _rect(s, 0, 0, ANCHO, 0.85, rp.COLOR_CAB)
        _texto(s, 0.4, 0.14, 9.6, 0.6, titulo, 24, "FFFFFF", True)
        _texto(s, 9.6, 0.14, 3.35, 0.32, f"Caso {self.esc.nombre_visible}", 11, "FFFFFF", True, PP_ALIGN.RIGHT)
        _texto(s, 9.6, 0.44, 3.35, 0.3, f"Fecha de Corrida: {self.fecha}", 9, "C9D6E3", alinear=PP_ALIGN.RIGHT)
        if subtitulo:
            _texto(s, 0.4, 0.92, 12.5, 0.34, subtitulo, 11, rp.COLOR_SUAVE)
        _texto(s, 8.4, 7.12, 4.55, 0.28, f"{config.NOMBRE_PROGRAMA} {config.VERSION_CORTA}", 9, rp.COLOR_SUAVE,
               alinear=PP_ALIGN.RIGHT)
        _texto(s, 0.4, 7.12, 2.0, 0.28, str(len(self.prs.slides)), 9, rp.COLOR_SUAVE)
        return s

    def titulo_bloque(self, s, x, y, w, texto):
        _texto(s, x, y, w, 0.3, texto, 12, rp.COLOR_TEXTO, True)

    def kpis(self, s, datos, top=1.3, izq=0.4, ancho=12.53):
        """Fila de indicadores: [(titulo, valor, detalle)]."""
        n = len(datos)
        sep = 0.15
        w = (ancho - sep * (n - 1)) / n
        for i, (titulo, valor, det) in enumerate(datos):
            x = izq + i * (w + sep)
            _rect(s, x, top, w, 0.82, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, borde="D0D7DE").adjustments[0] = 0.08
            _rect(s, x, top + 0.08, 0.05, 0.66, SERIES[i % len(SERIES)])
            _texto(s, x + 0.12, top + 0.04, w - 0.2, 0.28, titulo, 9, rp.COLOR_SUAVE)
            _texto(s, x + 0.12, top + 0.3, w - 0.2, 0.42, valor, 18, rp.COLOR_TEXTO, True)
            if det:
                _texto(s, x + w * 0.52, top + 0.4, w * 0.46, 0.3, det, 9, rp.COLOR_SUAVE, alinear=PP_ALIGN.RIGHT)
        return top + 0.95

    # ---- tablas -----------------------------------------------------------------------------------------------
    def tabla(self, slide, columnas, filas, totales=(), top=1.4, ancho=12.53, izq=0.4, alto_max=5.5,
              grupo=None, subtotales=()):
        """Tabla nativa; la altura de fila y la fuente se ajustan para que quepa en 'alto_max'.
        'grupo': índice de la columna que define el color de grupo."""
        n_f, n_c = len(filas) + 1, len(columnas)
        textos = [[_formato(tipo, f[j]) for j, (_k, _t, tipo) in enumerate(columnas)] for f in filas]
        # Fuente según el alto disponible; luego se reduce hasta que ninguna celda parta palabras o números
        tam = max(7, min(11 if n_c <= 8 else 10, int(min(0.34, alto_max / (n_f + 0.6)) * 36)))
        while True:
            cw = tam / 72 * 0.56                      # ancho medio de un carácter (pulgadas)
            nat = []
            for j, (_k, t, _tipo) in enumerate(columnas):
                palabra = max((len(w) for w in str(t).split()), default=1)
                dato = max([len(x[j]) for x in textos] + [0])
                nat.append(max(palabra * cw * 1.1, dato * cw) + 0.17)
            if sum(nat) <= ancho or tam <= 7:
                break
            tam -= 1
        extra = ancho - sum(nat)
        anchos = [w + extra * w / sum(nat) for w in nat]
        lineas = max(int(len(str(t)) * cw * 1.1 / max(0.2, a - 0.17)) + 1 for (_k, t, _ty), a in zip(columnas, anchos))
        h_cab = max(0.3, min(lineas, 3) * tam / 72 * 1.25 + 0.08)
        alto_fila = max(0.2, min(0.34, (alto_max - h_cab) / max(1, len(filas))))
        forma = slide.shapes.add_table(n_f, n_c, Inches(izq), Inches(top), Inches(ancho),
                                       Inches(h_cab + alto_fila * len(filas)))
        tabla = forma.table
        tabla.first_row = True
        for j, a in enumerate(anchos):
            tabla.columns[j].width = Inches(a)
        tabla.rows[0].height = Inches(h_cab)
        for i in range(1, n_f):
            tabla.rows[i].height = Inches(alto_fila)
        for j, (_k, titulo, tipo) in enumerate(columnas):
            self._celda(tabla.cell(0, j), titulo, tam, True, "FFFFFF", rp.COLOR_CAB,
                        PP_ALIGN.LEFT if tipo == "texto" else PP_ALIGN.CENTER)
        g, previo = -1, object()
        for i, fila in enumerate(filas, start=1):
            es_total, es_sub = (i - 1) in totales, (i - 1) in subtotales
            if grupo is not None and not es_total:
                if fila[grupo] != previo:
                    g, previo = g + 1, fila[grupo]
                fondo = GRUPOS[g % len(GRUPOS)]
            else:
                fondo = rp.COLOR_TOTAL if es_total else (rp.COLOR_ALT if i % 2 == 0 else "FFFFFF")
            for j, (_k, _t, tipo) in enumerate(columnas):
                self._celda(tabla.cell(i, j), textos[i - 1][j], tam, es_total or es_sub, rp.COLOR_TEXTO, fondo,
                            PP_ALIGN.LEFT if tipo == "texto" else PP_ALIGN.RIGHT)
        return top + h_cab + alto_fila * len(filas)

    @staticmethod
    def _celda(celda, texto, tam, negrita, color, fondo, alinear):
        celda.fill.solid()
        celda.fill.fore_color.rgb = _rgb(fondo)
        celda.vertical_anchor = MSO_ANCHOR.MIDDLE
        celda.margin_left = celda.margin_right = Inches(0.06)
        celda.margin_top = celda.margin_bottom = Inches(0.01)
        tf = celda.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = alinear
        r = p.add_run()
        r.text = texto
        r.font.name, r.font.size, r.font.bold = FUENTE, Pt(tam), negrita
        r.font.color.rgb = _rgb(color)

    # ---- gráficos nativos -------------------------------------------------------------------------------------
    def grafico(self, s, tipo, x, y, w, h, categorias, series, num="#,##0", etiquetas=True, leyenda=True,
                tam_etq=9, colores_puntos=None, titulo=None, invertir=None, max_eje=None):
        """series: [(nombre, valores, color|None|'ninguno')]. Devuelve el gráfico (editable en PowerPoint)."""
        cd = CategoryChartData()
        cd.categories = [str(c) for c in categorias]
        for nombre, valores, _c in series:
            cd.add_series(nombre, [None if v is None else round(float(v), 6) for v in valores])
        ch = s.shapes.add_chart(tipo, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
        ch.font.name, ch.font.size = FUENTE, Pt(9)
        ch.has_title = bool(titulo)
        if titulo:
            ch.chart_title.text_frame.text = titulo
            ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
            ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
        visibles = [k for k, (_n, _v, c) in enumerate(series) if c != "ninguno"]
        ch.has_legend = leyenda and len(visibles) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(9)
        ch.value_axis.has_major_gridlines = True
        ch.value_axis.major_gridlines.format.line.color.rgb = _rgb("E5E8EB")
        ch.value_axis.tick_labels.number_format = num.replace("0.0%", "0%")
        ch.value_axis.tick_labels.number_format_is_linked = False
        ch.value_axis.format.line.fill.background()
        ch.value_axis.tick_labels.font.size = Pt(9)
        ch.category_axis.tick_labels.font.size = Pt(9 if len(categorias) <= 30 else 7)
        if max_eje is not None:
            ch.value_axis.maximum_scale = max_eje
        horizontal = tipo in (XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.BAR_STACKED, XL_CHART_TYPE.BAR_STACKED_100)
        if invertir if invertir is not None else horizontal:
            ch.category_axis.reverse_order = True       # mismo orden que las tablas (de arriba hacia abajo)
        apilado = tipo in (XL_CHART_TYPE.BAR_STACKED, XL_CHART_TYPE.BAR_STACKED_100, XL_CHART_TYPE.COLUMN_STACKED,
                           XL_CHART_TYPE.COLUMN_STACKED_100)
        fmt_etq = num
        if apilado and etiquetas:
            # los segmentos pequeños no llevan etiqueta (no se enciman); el dato sigue en el gráfico
            if tipo in (XL_CHART_TYPE.BAR_STACKED_100, XL_CHART_TYPE.COLUMN_STACKED_100):
                umbral = 0.035
            else:
                sumas = [sum(float(v or 0) for (_n, vs, c) in series if c != "ninguno" for v in [vs[i]])
                         for i in range(len(categorias))]
                umbral = 0.06 * max(sumas + [0])
            if umbral > 0:
                fmt_etq = f'[<{umbral:.6g}]"";{num}'
        for k, (serie, (_n, _v, color)) in enumerate(zip(ch.series, series)):
            if color == "ninguno":
                serie.format.fill.background()
                serie.format.line.fill.background()
                continue
            serie.format.fill.solid()
            serie.format.fill.fore_color.rgb = _rgb(color or SERIES[k % len(SERIES)])
            if etiquetas:
                dl = serie.data_labels
                dl.show_value = True
                dl.number_format = fmt_etq
                dl.number_format_is_linked = False
                dl.font.size = Pt(tam_etq)
                if apilado:
                    dl.position = XL_LABEL_POSITION.CENTER
                    dl.font.color.rgb = _rgb("FFFFFF")
                else:
                    dl.position = XL_LABEL_POSITION.OUTSIDE_END
        if colores_puntos:
            serie = ch.series[-1]
            for i, c in enumerate(colores_puntos):
                p = serie.points[i]
                p.format.fill.solid()
                p.format.fill.fore_color.rgb = _rgb(c)
        try:
            ch.plots[0].gap_width = 60
            if tipo in (XL_CHART_TYPE.BAR_STACKED, XL_CHART_TYPE.BAR_STACKED_100, XL_CHART_TYPE.COLUMN_STACKED,
                        XL_CHART_TYPE.COLUMN_STACKED_100):
                ch.plots[0].overlap = 100
        except Exception:  # noqa
            pass
        return ch

    # ---- láminas ------------------------------------------------------------------------------------------------
    def resumen(self):
        res = self.res
        s = self.lamina("Resumen", f"Indicadores por clase y flota  ·  Réplica: {res.get('replica', '')}")
        clases = rp.resumen_clases(res)
        datos = []
        for c in clases:
            met = f"{c['metrica']:,.1f} Mt" if c["material"] else f"{c['metrica']:,.0f} taladros"
            datos.append((c["clase"], f"{c['n']:,} und", f"Disp. {fmt.porcentaje(c['disp'])} · "
                                                          f"Util. {fmt.porcentaje(c['util'])} · {met}"))
        y = self.kpis(s, [(t, v, "") for t, v, _d in datos]) if datos else 1.3
        for i, (_t, _v, d) in enumerate(datos):
            w = (12.53 - 0.15 * (len(datos) - 1)) / len(datos)
            _texto(s, 0.4 + i * (w + 0.15) + w * 0.3, 1.72, w * 0.68, 0.3, d, 8.5, rp.COLOR_SUAVE, alinear=PP_ALIGN.RIGHT)
        filas_r = rp.resumen_flotas(res)
        cols = [("c", "Clase", "texto"), ("f", "Flota", "texto"), ("n", "N° de Equipos", "entero"),
                ("h", "Horas Productivas (h)", "entero"), ("d", "Disponibilidad (%)", "pct"),
                ("u", "Utilización (%)", "pct"), ("t", "Material (Mt) / N° Taladros", "texto")]
        filas = [[r["clase"], r["flota"], r["n"], r["hp"], r["disp"], r["util"],
                  fmt.decimal1(r["metrica"]) if r["material"] else fmt.entero(r["metrica"])] for r in filas_r]
        self.tabla(s, cols, filas, top=y + 0.05, ancho=7.0, alto_max=7.0 - y - 0.1, grupo=0)
        if filas_r:
            self.grafico(s, XL_CHART_TYPE.BAR_CLUSTERED, 7.6, y, 5.33, 7.05 - y, [r["flota"] for r in filas_r],
                         [("Disponibilidad (%)", [r["disp"] for r in filas_r], SERIES[0]),
                          ("Utilización (%)", [r["util"] for r in filas_r], SERIES[1])],
                         num="0.0%", tam_etq=8, titulo="Disponibilidad y Utilización por Flota", max_eje=1.0,
                         etiquetas=len(filas_r) <= 14)

    def estados(self):
        """Análisis de Estados: por clase, una lámina Resumen (KPIs, distribución, Por Flota y Por Tipo)
        y una lámina Por ID (gráfico de todos los equipos + resumen por flota)."""
        for clase in modelo.ORDEN_CLASES:
            r = self.res["clases"].get(clase)
            if not r or not r["flota"]:
                continue
            info = modelo.CLASES[clase]
            t = r["total"]
            s = self.lamina(f"Análisis de Estados – {info['plural']}", f"Resumen  ·  Réplica: {r.get('replica', '')}")
            met = (f"{t['metrica_total']:,.1f}" if info["metrica"] == "material" else f"{t['metrica_total']:,.0f}")
            y = self.kpis(s, [(f"N° de {info['plural']}", f"{r['n_equipos']:,}", ""),
                              ("Disponibilidad", fmt.porcentaje(t["disponibilidad"]), ""),
                              ("Utilización", fmt.porcentaje(t["utilizacion"]), ""),
                              ("Horas Productivas (h)", fmt.entero(t["horas_productivas"]), ""),
                              (info["metrica_total"], met, "")])
            flotas = r["flota"]
            alto_g = min(2.6, 0.75 + 0.32 * len(flotas))
            self.grafico(s, XL_CHART_TYPE.BAR_STACKED_100, 0.4, y, 12.53, alto_g, [f["nombre"] for f in flotas],
                         [(modelo.NOMBRE_ESTADO[k], [(f.get(k.lower(), 0) / f["tc"]) if f.get("tc") else 0
                                                     for f in flotas], rp.COLOR_ESTADO[k]) for k in modelo.CLAVES],
                         num="0.0%", tam_etq=8, titulo="Distribución del Tiempo Calendario (%)")
            y += alto_g + 0.1
            tablas = {t_["titulo"]: t_ for t_ in rp.tablas_clase(r, clase)}
            alto = 7.05 - y - 0.3
            for k, nombre in enumerate(("Por Flota", "Por Tipo")):
                tb = tablas[nombre]
                x = 0.4 + k * 6.34
                self.titulo_bloque(s, x, y, 6.1, nombre)
                self.tabla(s, tb["columnas"], tb["filas"], tb["totales"], top=y + 0.3, ancho=6.19, izq=x, alto_max=alto)
            self._por_id(r, clase)

    def _por_id(self, r, clase):
        info = modelo.CLASES[clase]
        ids = r["id"]
        if not ids:
            return
        s = self.lamina(f"Análisis de Estados – {info['plural']} – Por ID",
                        f"{len(ids)} equipos  ·  la tabla resume cada flota (mínimo, promedio y máximo por equipo)")
        self.grafico(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 0.4, 1.3, 12.53, 2.9, [str(e["nombre"]) for e in ids],
                     [("Disponibilidad (%)", [e["disponibilidad"] for e in ids], SERIES[0]),
                      ("Utilización (%)", [e["utilizacion"] for e in ids], SERIES[1])],
                     num="0%", etiquetas=len(ids) <= 16, titulo="Disponibilidad y Utilización por ID", max_eje=1.0)
        filas = []
        for f in r["flota"]:
            del_f = [e for e in ids if e["flota"] == f["nombre"]]
            if not del_f:
                continue
            d = [e["disponibilidad"] for e in del_f if e.get("disponibilidad") is not None]
            u = [e["utilizacion"] for e in del_f if e.get("utilizacion") is not None]
            m = [e["metrica_total"] for e in del_f]
            filas.append([f.get("tipo", ""), f["nombre"], len(del_f),
                          min(d) if d else None, sum(d) / len(d) if d else None, max(d) if d else None,
                          min(u) if u else None, sum(u) / len(u) if u else None, max(u) if u else None,
                          sum(m) / len(m) if m else None])
        met = "dec1" if info["metrica"] == "material" else "entero"
        cols = [("t", "Tipo", "texto"), ("f", "Flota", "texto"), ("n", "N° de IDs", "entero"),
                ("d0", "Disp. Mínima", "pct"), ("d1", "Disp. Promedio", "pct"), ("d2", "Disp. Máxima", "pct"),
                ("u0", "Util. Mínima", "pct"), ("u1", "Util. Promedio", "pct"), ("u2", "Util. Máxima", "pct"),
                ("m", f"{info['metrica_id']} (Promedio por ID)", met)]
        self.tabla(s, cols, filas, top=4.35, alto_max=2.65, grupo=0)

    def tiempos(self):
        """Análisis de Tiempos: una lámina por clase con el gráfico apilado (h/día-eq) e indicadores."""
        tiempos = self.res.get("tiempos") or {}
        for clase in modelo.ORDEN_CLASES:
            t = tiempos.get(clase)
            if not t or not t.get("ids"):
                continue
            info = modelo.CLASES[clase]
            bloques = []
            for flota in modelo.flotas_de(t):
                ids = [e for e in t["ids"] if e["flota"] == flota]
                b = modelo.bloque_tiempos(t, ids, "dia", self.esc.dias)
                ba = modelo.bloque_tiempos(t, ids, "anio", self.esc.dias)
                if b:
                    bloques.append((flota, b, ba))
            if not bloques:
                continue
            s = self.lamina(f"Análisis de Tiempos – {info['plural']}",
                            f"Horas por equipo y día (h/día-eq) y por periodo (h/periodo-eq)  ·  "
                            f"Réplica: {self.res.get('replica', '')}")
            paso = modelo.paso_eje("dia", self.esc.dias)
            ch = self.grafico(s, XL_CHART_TYPE.BAR_STACKED, 0.4, 1.3, 7.0, 5.7, [f for f, _b, _a in bloques],
                              [(modelo.FILA_TIEMPOS[k], [b["por_estado"][k] for _f, b, _a in bloques], rp.COLOR_ESTADO[k])
                               for k in modelo.ORDEN_TIEMPOS], num="0.0", tam_etq=8,
                              titulo="Distribución por Estado (h/día-eq)")
            ch.value_axis.major_unit = paso
            ch.value_axis.minimum_scale = 0
            if max(b["total"] for _f, b, _a in bloques) <= paso * 6 * 1.0005:
                ch.value_axis.maximum_scale = paso * 6          # 24 h/día-eq
            cols = [("f", "Flota", "texto"), ("n", "Und", "entero"), ("h", "Horas Prod. (h/día-eq)", "dec2"),
                    ("d", "Disponibilidad", "pct"), ("u", "Utilización", "pct"), ("t", "h/día-eq", "dec2"),
                    ("a", "h/periodo-eq", "dec1")]
            filas = [[f, b["n"], b["horas_productivas"], b["disponibilidad"], b["utilizacion"], b["total"],
                      a["total"] if a else None] for f, b, a in bloques]
            self.titulo_bloque(s, 7.6, 1.3, 5.3, "Indicadores por Flota")
            self.tabla(s, cols, filas, top=1.62, ancho=5.33, izq=7.6, alto_max=5.3)

    def cargas(self):
        """Orígenes y Destinos: cascadas por tipo de origen y de destino + tablas (sin filtros)."""
        import registro
        c = self.res.get("cargas")
        if not c:
            return
        s = self.lamina("Orígenes y Destinos", f"Material en millones de toneladas (Mt)  ·  Réplica: "
                                               f"{self.res.get('replica', '')}")
        for k, (titulo, datos, tipos, col) in enumerate((
                ("Reporte por Tipo de Origen", registro.por_tipo_origen(c, self.esc.origenes), registro.TIPOS_ORIGEN,
                 "Tipo de Origen"),
                ("Reporte por Tipo de Destino", registro.por_tipo_destino(c, self.esc.destinos), registro.TIPOS_DESTINO,
                 "Tipo de Destino"))):
            tipos = [t for t in tipos if datos["tipos"].get(t, 0) > 0 or t != "N/A"]
            vals = [datos["tipos"].get(t, 0.0) / 1e6 for t in tipos]
            base, acum = [], 0.0
            for v in vals:
                base.append(acum)
                acum += v
            x = 0.4 + k * 6.34
            self.titulo_bloque(s, x, 1.3, 6.1, titulo)
            self.grafico(s, XL_CHART_TYPE.COLUMN_STACKED, x, 1.62, 6.19, 3.5, tipos + ["Total"],
                         [("Base", base + [0.0], "ninguno"), ("Material (Mt)", vals + [datos["total"] / 1e6], None)],
                         num="#,##0.0", leyenda=False, tam_etq=10,
                         colores_puntos=[COLOR_TIPO.get(t, SERIES[0]) for t in tipos + ["Total"]], invertir=False)
            total = datos["total"] or 1.0
            filas = [[t, datos["tipos"].get(t, 0.0) / 1e6, datos["tipos"].get(t, 0.0) / total] for t in tipos]
            filas.append(["Total", datos["total"] / 1e6, 1.0])
            self.tabla(s, [("t", col, "texto"), ("m", "Material (Mt)", "dec2"), ("p", "Participación (%)", "pct")],
                       filas, totales=[len(filas) - 1], top=5.3, ancho=6.19, izq=x, alto_max=1.7)

    def tiempos_ciclo(self):
        """Tiempos y Métricas ▸ Tiempos Ciclo: promedio por flota de pala y tipo de camión de cada tiempo y
        Descarga Plan (las láminas resumen el promedio; mínimo y máximo están en el reporte Excel)."""
        import registro
        c = self.res.get("cargas")
        tp = (self.esc.clases.get("palas") or {}).get("tipos") or {}
        tc = (self.esc.clases.get("camiones") or {}).get("tipos") or {}
        if not c or not tp or not tc or not registro.tiene_estadisticas(c):
            return
        cfg_p, cfg_c = self.esc.clases.get("palas"), self.esc.clases.get("camiones")
        s = self.lamina("Tiempos y Métricas – Tiempos Ciclo",
                        f"Promedio por registro de carga (min)  ·  Réplica: {self.res.get('replica', '')}")
        medidas = [(m, t.replace("Tiempo de ", "")) for m, t, _u in registro.MEDIDAS_TIEMPO]
        tablas = {m: registro.tabla_estadistica(c, m, tp, tc, cfg_p, cfg_c) for m, _t in medidas}
        claves = list(dict.fromkeys((f["pala"], f["camion"]) for m, _t in medidas for f in tablas[m]))
        prom = {m: {(f["pala"], f["camion"]): f for f in tablas[m]} for m, _t in medidas}
        filas = [[p, t] + [(prom[m].get((p, t)) or {}).get("prom") for m, _t in medidas] for p, t in claves]
        self.titulo_bloque(s, 0.4, 1.3, 12.5, "Tiempos Promedio por Flota de Pala y Tipo de Camión (min)")
        y = self.tabla(s, [("p", "Flota Equipo Pala", "texto"), ("c", "Tipo Camiones", "texto")] +
                       [(m, t, "dec2") for m, t in medidas], filas, top=1.62, ancho=7.3, alto_max=2.6, grupo=0)
        self.grafico(s, XL_CHART_TYPE.BAR_STACKED, 7.9, 1.3, 5.03, max(2.9, y - 1.3),
                     [f"{p} – {t}" for p, t in claves],
                     [(t, [(prom[m].get(k) or {}).get("prom") for k in claves], SERIES[j]) for j, (m, t) in
                      enumerate(medidas) if m in ("cg", "cu", "hang", "cc")], num="0.0", tam_etq=8,
                     titulo="Ciclo en Carga (min)", etiquetas=len(claves) <= 8)
        desc = registro.tabla_descarga(c, self.esc.destinos, tc)
        if desc:
            y = max(y, 4.3) + 0.15
            self.titulo_bloque(s, 0.4, y, 7.0, "Tiempo de Descarga Plan (min)")
            self.tabla(s, [("d", "Tipo de Destino", "texto"), ("c", "Tipo de Camión", "texto"),
                           ("q", "Prom. Cola en Descarga", "dec2"), ("x", "Prom. Descarga", "dec2"),
                           ("t", "Descarga Total", "dec2")],
                       [[f["tipo_destino"], f["camion"], f["cola"], f["descarga"], f["total"]] for f in desc],
                       top=y + 0.32, ancho=7.3, alto_max=7.0 - y - 0.35, grupo=0)
            self.grafico(s, XL_CHART_TYPE.BAR_STACKED, 7.9, y, 5.03, 7.05 - y,
                         [f"{f['tipo_destino']} – {f['camion']}" for f in desc],
                         [("Cola en Descarga", [f["cola"] for f in desc], rp.COLOR_ESTADO["DPNP"]),
                          ("Descarga", [f["descarga"] for f in desc], rp.COLOR_ESTADO["TP"])],
                         num="0.00", tam_etq=8, titulo="Descarga Plan (min)")

    def metricas(self):
        """Tiempos y Métricas ▸ Métricas: (1) Registro de Ciclos, N° de Pases y Carga por Camión;
        (2) Productividad de Palas y Camiones y N° Taladros."""
        import registro
        res = self.res
        c = res.get("cargas")
        tp = (self.esc.clases.get("palas") or {}).get("tipos") or {}
        tc = (self.esc.clases.get("camiones") or {}).get("tipos") or {}
        cfg_p, cfg_c = self.esc.clases.get("palas"), self.esc.clases.get("camiones")
        if c and tp and tc and registro.tiene_estadisticas(c):
            s = self.lamina("Tiempos y Métricas – Métricas de Ciclo",
                            f"Registro de Ciclos, N° de Pases y Carga por Camión  ·  Réplica: {res.get('replica', '')}")
            r = registro.tabla_registro(c, tp, tc, cfg_p, cfg_c)
            palas, cams = r["palas"], r["camiones"]
            total = sum(r["n"].get((p, k), 0) for p in palas for k in cams) or 1
            filas = [[p] + [r["n"].get((p, k), 0) for k in cams] + [sum(r["n"].get((p, k), 0) for k in cams),
                                                                     sum(r["n"].get((p, k), 0) for k in cams) / total,
                                                                     sum(r["t"].get((p, k), 0) for k in cams) / 1e6]
                     for p in palas]
            filas.append(["Total"] + [sum(r["n"].get((p, k), 0) for p in palas) for k in cams] +
                         [total, 1.0, sum(r["t"].values()) / 1e6])
            self.titulo_bloque(s, 0.4, 1.3, 6.2, "Registro de Ciclos")
            self.tabla(s, [("p", "Flota Equipo Pala", "texto")] + [(k, k, "entero") for k in cams] +
                       [("n", "N° Ciclos", "entero"), ("pc", "Ciclos (%)", "pct"), ("t", "Tonelaje (Mt)", "dec2")],
                       filas, totales=[len(filas) - 1], top=1.62, ancho=6.19, alto_max=2.4)
            if palas:
                self.grafico(s, XL_CHART_TYPE.BAR_STACKED, 6.74, 1.3, 6.19, 2.75, palas,
                             [(k, [r["n"].get((p, k), 0) for p in palas], SERIES[0] if k == "UltraClass" else SERIES[2])
                              for k in cams], num="#,##0", tam_etq=8, titulo="Registro de Ciclos – N° Ciclos")
            for k, (med, titulo, num) in enumerate((("ps", "N° de Pases", "0.0"), ("t", "Carga por Camión (t)", "#,##0"))):
                filas = registro.tabla_estadistica(c, med, tp, tc, cfg_p, cfg_c)
                if not filas:
                    continue
                x = 0.4 + k * 6.34
                self.grafico(s, XL_CHART_TYPE.BAR_CLUSTERED, x, 4.2, 6.19, 2.85,
                             [f"{f['pala']} – {f['camion']}" for f in filas],
                             [("Mínimo", [f["min"] for f in filas], "8C959F"),
                              ("Promedio", [f["prom"] for f in filas], "12355B"),
                              ("Máximo", [f["max"] for f in filas], "3B8ED8")], num=num, tam_etq=7,
                             titulo=titulo, etiquetas=len(filas) <= 6)
        clases = res.get("clases") or {}
        if not (clases.get("palas") or clases.get("camiones") or clases.get("perforadoras")):
            return
        s = self.lamina("Tiempos y Métricas – Productividad y N° Taladros",
                        f"Productividad por flota (t/h) y taladros perforados  ·  Réplica: {res.get('replica', '')}")

        def prod(f):
            return f["metrica_total"] * 1e6 / f["horas_productivas"] if f.get("horas_productivas") else None
        k = 0
        for clase, titulo in (("palas", "Productividad Palas (t/h)"), ("camiones", "Productividad Camiones (t/h)")):
            r = clases.get(clase)
            if not r:
                continue
            x = 0.4 + k * 6.34
            self.grafico(s, XL_CHART_TYPE.BAR_CLUSTERED, x, 1.3, 6.19, 2.15, [f["nombre"] for f in r["flota"]],
                         [("Productividad (t/h)", [prod(f) for f in r["flota"]], SERIES[k])], num="#,##0",
                         leyenda=False, titulo=f"{titulo} – Por Flota")
            filas = [[f.get("tipo", ""), f["nombre"], f["horas_productivas"], f["metrica_total"], prod(f)]
                     for f in r["flota"]]
            self.tabla(s, [("t", "Tipo", "texto"), ("f", "Flota", "texto"), ("h", "Horas Productivas (h)", "entero"),
                           ("m", "Material (Mt)", "dec2"), ("p", "Productividad (t/h)", "entero")], filas, top=3.5,
                       ancho=6.19, izq=x, alto_max=1.55, grupo=0)
            k += 1
        perf = clases.get("perforadoras")
        if perf:
            cat = (self.esc.clases.get("perforadoras") or {}).get("energia") or {}
            filas = [[f.get("tipo") or modelo.SIN_ASIGNAR, f["nombre"], cat.get(f["nombre"], ""), f["n_equipos"],
                      f["metrica_total"]] for f in perf["flota"]]
            filas.append(["Total", "", "", perf["n_equipos"], sum(f["metrica_total"] for f in perf["flota"])])
            self.titulo_bloque(s, 0.4, 5.15, 6.0, "N° Taladros")
            self.tabla(s, [("t", "Tipo de Perforadora", "texto"), ("f", "Flota", "texto"),
                           ("c", modelo.CATEGORIA, "texto"), ("n", "N° de Perforadoras", "entero"),
                           ("x", "N° Taladros", "entero")], filas, totales=[len(filas) - 1], top=5.45, ancho=12.53,
                       alto_max=1.6)

    def plan_mina(self):
        """Plan de Mina mensual (sin filtros): indicadores, tonelaje por mes y flota de pala y resumen por fase."""
        import registro
        from vista_perfil import orden_flotas
        c = self.res.get("cargas")
        if not c or not c.get("perfil"):
            return
        t = registro.tabla_perfil(c, self.esc.anio or 2000, "mes", self.esc.origenes, self.esc.destinos, {},
                                  orden_flotas(self.esc))
        if not t["filas"]:
            return
        s = self.lamina("Plan de Mina", f"Tonelaje mensual (t) por flota de pala  ·  Año {self.esc.anio}  ·  "
                                        f"Réplica: {self.res.get('replica', '')}")
        per, tot = t["periodos"], t["totales"]
        con = [i for i in range(len(tot)) if tot[i] > 0] or [0]
        i_max, i_min = max(con, key=lambda i: tot[i]), min(con, key=lambda i: tot[i])
        y = self.kpis(s, [("Tonelaje Movido Total (t)", fmt.entero(t["total"]), ""),
                          ("Tonelaje Promedio Mensual (t)", fmt.entero(t["total"] / len(con)), ""),
                          ("Tonelaje Máximo Mensual (t)", fmt.entero(tot[i_max]), per[i_max]),
                          ("Tonelaje Mínimo Mensual (t)", fmt.entero(tot[i_min]), per[i_min])])
        orden = orden_flotas(self.esc)
        flotas = sorted(dict.fromkeys(f["flota"] for f in t["filas"]),
                        key=lambda x: orden.index(x) if x in orden else len(orden))
        self.grafico(s, XL_CHART_TYPE.COLUMN_STACKED, 0.4, y, 8.2, 7.05 - y, per,
                     [(fl, [sum(f["valores"][i] for f in t["filas"] if f["flota"] == fl) / 1e3 for i in range(len(per))],
                       SERIES[j % len(SERIES)]) for j, fl in enumerate(flotas)], num="#,##0", tam_etq=7,
                     titulo="Tonelaje por Mes y Flota Pala (kt)", etiquetas=len(per) <= 12 and len(flotas) <= 4)
        fases = list(dict.fromkeys(f["fase"] for f in t["filas"]))
        filas = []
        for fase in fases:
            v = sum(f["total"] for f in t["filas"] if f["fase"] == fase)
            filas.append([fase, v / 1e6, v / (t["total"] or 1.0)])
        filas.sort(key=lambda x: -x[1])
        if len(filas) > MAX_FILAS - 6:                      # se resume: las principales y «Otras»
            resto = filas[MAX_FILAS - 7:]
            filas = filas[:MAX_FILAS - 7] + [[f"Otras ({len(resto)})", sum(x[1] for x in resto), sum(x[2] for x in resto)]]
        filas.append(["Total", t["total"] / 1e6, 1.0])
        self.titulo_bloque(s, 8.8, y, 4.1, "Tonelaje por Fase")
        self.tabla(s, [("f", "Fase", "texto"), ("m", "Tonelaje (Mt)", "dec2"), ("p", "Participación (%)", "pct")],
                   filas, totales=[len(filas) - 1], top=y + 0.32, ancho=4.13, izq=8.8, alto_max=7.0 - y - 0.35)

    def plan_vs_sim(self):
        """Plan vs Simulación: una lámina por sección del Vector Plan (tabla + gráfico nativo)."""
        import plan
        esc = self.esc
        reps = plan.replicas_resultados(esc)
        if not esc.plan or not reps:
            return
        por = self.res["por_replica"]
        multi = len(reps) > 1
        nombre_plan = esc.plan.get("nombre", "Plan")
        secciones = []
        for f in esc.plan["filas"]:
            if f["tipo"] == "seccion":
                secciones.append((f["texto"].strip(), []))
            elif secciones:
                secciones[-1][1].append(f)
            else:
                secciones.append(("Vector Plan", [f]))
        for titulo, items in secciones:
            if not items:
                continue
            s = self.lamina(f"Plan vs Simulación – {titulo[:60]}", f"{nombre_plan}  vs  {esc.nombre_visible}")
            cols = [("p", "Parámetro", "texto"), ("u", "Unidad", "texto"), ("a", nombre_plan, "texto")]
            for r in reps:
                cols += [(f"s{r}", f"{esc.nombre_visible} – Réplica {r}" if multi else esc.nombre_visible, "texto"),
                         (f"v{r}", f"Variación Réplica {r}" if multi else "Variación", "texto")]
            filas, sims_all = [], []
            for f in items:
                sims = [plan.valor(esc, f.get("vinculo"), por[r]) for r in reps]
                sims_all.append(sims)
                fila = [f["texto"].strip(), f["unidad"], plan.texto_valor(f, f["valor"])]
                for v in sims:
                    fila += [plan.texto_valor(f, v),
                             plan.texto_variacion(f, f.get("variacion"), plan.variacion(f.get("variacion"), f["valor"], v))]
                filas.append(fila)
            ancho_t = 7.4 if len(items) <= 22 else 12.53
            self.tabla(s, cols, filas, top=1.35, ancho=ancho_t, alto_max=5.65)
            if ancho_t < 12:
                pct = plan.es_porcentaje(items[0])
                factor = 100.0 if pct else 1.0
                validos = [(f, sv) for f, sv in zip(items, sims_all) if f["valor"] is not None or any(v is not None for v in sv)]
                if validos:
                    series = [(nombre_plan, [(f["valor"] or 0.0) * factor for f, _s in validos], "8C959F")]
                    for j, r in enumerate(reps):
                        series.append((f"Réplica {r}" if multi else esc.nombre_visible,
                                       [(sv[j] or 0.0) * factor for _f, sv in validos], SERIES[j % len(SERIES)]))
                    self.grafico(s, XL_CHART_TYPE.BAR_CLUSTERED, 8.0, 1.35, 4.93, 5.65,
                                 [f["texto"].strip()[:40] for f, _s in validos], series,
                                 num="0.0" if pct or plan.decimales(validos[0][0]) else "#,##0", tam_etq=7,
                                 etiquetas=len(validos) <= 8)

    def pivots(self):
        """Tablas Dinámicas (configuración guardada o base): tabla y gráfico por clase; las tablas que no
        caben se resumen por el primer nivel de filas."""
        import copy
        import pivot
        from vista_pivot_estados import config_inicial
        reps = [str(r) for r in self.res.get("replicas", []) if r != "Todas"]
        for clase in modelo.ORDEN_CLASES:
            if clase not in self.res.get("clases", {}):
                continue
            df = modelo.dataset_estados(self.res, clase, self.esc.clases.get(clase))
            if df.empty:
                continue
            cfg = dict(config_inicial(clase, reps))
            guardado = (self.esc.pivots.get(clase) or {}).get("config")
            if guardado:
                cfg.update(guardado)
            cfg["filtros"] = []
            try:
                tabla = pivot.calcular(df, copy.deepcopy(cfg), modelo.MEDIDAS_ESTADOS)
            except Exception:  # noqa
                continue
            if len(tabla["filas"]) > MAX_FILAS and len(cfg.get("filas", [])) > 1:
                cfg = dict(cfg, filas=cfg["filas"][:1])
                tabla = pivot.calcular(df, copy.deepcopy(cfg), modelo.MEDIDAS_ESTADOS)
            info = modelo.CLASES[clase]
            s = self.lamina(f"Tabla Dinámica – Análisis de Estados – {info['plural']}", self._resumen_cfg(cfg))
            self._tabla_pivot(s, tabla, f"Análisis de Estados – {info['plural']}")
        guardado = self.esc.pivot if isinstance(self.esc.pivot, dict) else None
        if guardado and guardado.get("resultado"):
            s = self.lamina("Tabla Dinámica – Registro de Cargas", self._resumen_cfg(guardado.get("config") or {}))
            self._tabla_pivot(s, guardado["resultado"], "Registro de Cargas")

    @staticmethod
    def _resumen_cfg(cfg):
        partes = []
        for clave, nombre in (("filas", "Filas"), ("columnas", "Columnas"), ("filtros", "Filtros")):
            if cfg.get(clave):
                partes.append(f"{nombre}: " + ", ".join(d["campo"] for d in cfg[clave]))
        return "  ·  ".join(partes)

    def _tabla_pivot(self, s, tabla, titulo):
        """Tabla del pivote (resumida si no cabe) y, si queda espacio, su gráfico nativo."""
        import pivot
        n_dim = tabla["n_dim"]
        formatos = tabla.get("formatos") or [None] * len(tabla["columnas"])
        cols = []
        for j, t in enumerate(tabla["columnas"]):
            if j < n_dim:
                cols.append((f"c{j}", t, "texto"))
                continue
            f = formatos[j] if j < len(formatos) else None
            tipo = "pct" if f in ("pct", "0.0%", "0.00%") else ("entero" if f in ("int", "#,##0") else "dec1")
            vals = [x[j] for x in tabla["filas"] if isinstance(x[j], float)]
            if tipo == "dec1" and ("(h)" in t or (vals and all(abs(v - round(v)) < 1e-9 for v in vals))):
                tipo = "entero"
            cols.append((f"c{j}", str(t).replace("\n", " · "), tipo))
        filas = list(tabla["filas"])
        totales = list(tabla["totales"])
        if len(filas) > MAX_FILAS:                      # se conservan las primeras filas y el total general
            filas = filas[:MAX_FILAS - 1] + [filas[i] for i in totales]
            totales = list(range(MAX_FILAS - 1, len(filas)))
        if len(cols) > 12:
            cols, filas = cols[:12], [f[:12] for f in filas]
        cuerpo = [i for i in range(len(filas)) if i not in totales]
        con_grafico = 0 < len(cuerpo) <= 14
        alto = min(5.65, 0.34 * (len(filas) + 2)) if con_grafico else 5.65
        y = self.tabla(s, cols, filas, totales, top=1.35, alto_max=alto, grupo=0 if n_dim > 1 else None)
        if not con_grafico or y > 5.0:
            return
        numericas = [j for j in range(n_dim, len(cols)) if pivot.TOTAL not in str(cols[j][1])]
        pct = [j for j in numericas if cols[j][2] == "pct"]
        sel = (pct or numericas[:1])[:6]
        if not sel:
            return
        cats = [" · ".join(str(x) for x in filas[i][:n_dim] if str(x)) for i in cuerpo]
        self.grafico(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 0.4, y + 0.15, 12.53, 7.05 - y - 0.15, cats,
                     [(cols[j][1], [filas[i][j] if isinstance(filas[i][j], (int, float)) else None for i in cuerpo],
                       SERIES[k % len(SERIES)]) for k, j in enumerate(sel)],
                     num="0.0%" if pct else "#,##0", tam_etq=8, titulo=titulo, etiquetas=len(cuerpo) * len(sel) <= 40)


def exportar_reporte(ruta, esc, progreso=None):
    if not esc.resultados or not esc.resultados.get("clases"):
        raise ValueError("No hay resultados para exportar.")
    rep = Reporte(esc)
    pasos = [("Portada", rep.portada), ("Resumen", rep.resumen), ("Análisis de Estados", rep.estados),
             ("Análisis de Tiempos", rep.tiempos), ("Orígenes y Destinos", rep.cargas),
             ("Tiempos Ciclo", rep.tiempos_ciclo), ("Métricas", rep.metricas), ("Plan de Mina", rep.plan_mina),
             ("Plan vs Simulación", rep.plan_vs_sim), ("Tablas Dinámicas", rep.pivots)]
    for i, (nombre, fn) in enumerate(pasos):
        if progreso:
            progreso(i / (len(pasos) + 1), f"Lámina «{nombre}»…")
        fn()
    if progreso:
        progreso(0.97, "Guardando la presentación…")
    rep.prs.save(ruta)
    return ruta
