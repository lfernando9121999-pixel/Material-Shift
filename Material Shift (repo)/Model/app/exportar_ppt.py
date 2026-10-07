"""PowerPoint con gráficos NATIVOS y editables (python-pptx), igual que SimA.

spec = {"charts": [chart, ...], "titulo": "...", "subtitulo": "..."}   (o un solo chart)
chart = {
  "title": "Material semanal por destino (t)", "subtitle": "Escenario · Desmonte",
  "kind": "line" | "column" | "bar",          # bar = barras horizontales
  "categories": [1, 2, ...],
  "series": [{"name": "Chw2a_1", "values": [...], "color": "#1f77b4", "dash": false}],
  "numberFormat": "#,##0"
}
Una lámina 16:9 por gráfico. La app copia la forma del gráfico al portapapeles con
PowerPoint (COM) para pegarla como gráfico editable, o guarda el .pptx.
"""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x0B, 0x25, 0x45)
GRIS = RGBColor(0x60, 0x70, 0x87)


def _rgb(hex_color, default="1F77B4"):
    h = (hex_color or default).lstrip("#")
    if len(h) != 6:
        h = default
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _agregar(slide, chart, x, y, w, h):
    kind = chart.get("kind", "line")
    tipo = {"line": XL_CHART_TYPE.LINE, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "bar": XL_CHART_TYPE.BAR_CLUSTERED}.get(kind, XL_CHART_TYPE.LINE)
    datos = CategoryChartData(number_format=chart.get("numberFormat", "#,##0"))
    datos.categories = [str(c) for c in chart.get("categories", [])]
    series = chart.get("series", [])
    for s in series:
        datos.add_series(str(s.get("name", "")), [_num(v) for v in s.get("values", [])])
    gf = slide.shapes.add_chart(tipo, x, y, w, h, datos)
    ch = gf.chart
    ch.font.size = Pt(10)
    ch.font.name = "Segoe UI"
    ch.has_legend = len(series) > 0
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(10)
    ch.has_title = False
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE8, 0xEF)
    va.tick_labels.number_format = chart.get("numberFormat", "#,##0")
    va.tick_labels.number_format_is_linked = False
    va.format.line.fill.background()
    ca = ch.category_axis
    ca.format.line.color.rgb = RGBColor(0xC9, 0xD2, 0xDD)
    ca.tick_labels.font.size = Pt(9)
    if kind == "line" and len(datos.categories) > 60:
        # Fewer labels on long daily axes so they stay readable.
        ca_el = ca._element
        from pptx.oxml.ns import qn
        from lxml import etree
        for tag, val in (("c:tickLblSkip", "30"), ("c:tickMarkSkip", "30")):
            el = etree.SubElement(ca_el, qn(tag))
            el.set("val", val)
    plot = ch.plots[0]
    if kind in ("column", "bar"):
        plot.gap_width = 60
        plot.overlap = -10 if len(series) > 1 else 0
    for i, s in enumerate(series):
        ser = plot.series[i]
        color = _rgb(s.get("color"))
        if kind == "line":
            ser.smooth = False
            ser.marker.style = None
            from pptx.enum.chart import XL_MARKER_STYLE
            ser.marker.style = XL_MARKER_STYLE.NONE
            ln = ser.format.line
            ln.color.rgb = color
            ln.width = Pt(s.get("width", 2) * 0.9)
            if s.get("dash"):
                ln.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = color
    return gf


def crear_pptx(spec, salida):
    charts = spec.get("charts") if isinstance(spec, dict) and "charts" in spec else [spec]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for chart in charts:
        slide = prs.slides.add_slide(blank)
        # Chart first, so it is Shapes.Item(1) (the app copies that shape).
        _agregar(slide, chart, Inches(0.5), Inches(1.25), Inches(12.33), Inches(5.85))
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12.33), Inches(0.5))
        p = tb.text_frame.paragraphs[0]
        p.text = chart.get("title", "")
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.font.name = "Segoe UI"
        if chart.get("subtitle"):
            p2 = tb.text_frame.add_paragraph()
            p2.text = chart["subtitle"]
            p2.font.size = Pt(12)
            p2.font.color.rgb = GRIS
            p2.font.name = "Segoe UI"
            p2.alignment = PP_ALIGN.LEFT
    salida = Path(salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    prs.save(salida)
    return salida
